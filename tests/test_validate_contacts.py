"""Tests for scripts/validate_contacts.py."""
import json
import sys
from pathlib import Path

import numpy as np
import pandas as pd
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
import validate_contacts as vc  # noqa: E402


def _events(n):
    return pd.DataFrame({"start_frame": np.arange(n) * 100, "end_frame": np.arange(n) * 100 + 10,
                         "duration_sec": np.linspace(0.1, 3.0, n)})


def test_tertiles_even():
    t = vc.assign_tertiles(pd.Series(np.linspace(0.1, 3, 30)))
    assert t.value_counts().to_dict() == {"short": 10, "medium": 10, "long": 10}


def test_sample_stratified_and_deterministic():
    df = _events(60)
    a = vc.stratified_sample(df, 15, 42)
    b = vc.stratified_sample(df, 15, 42)
    c = vc.stratified_sample(df, 15, 7)
    assert len(a) == 15
    assert a["duration_tertile"].value_counts().to_dict() == {"short": 5, "medium": 5, "long": 5}
    assert list(a.index) == list(b.index)
    assert list(a.index) != list(c.index)


def test_sample_takes_all_when_few():
    assert len(vc.stratified_sample(_events(5), 15, 1)) == 5


def test_sample_quota_when_uneven():
    assert len(vc.stratified_sample(_events(20), 19, 1)) == 19


def test_wilson():
    lo, hi = vc.wilson_ci(8, 10)
    assert lo == pytest.approx(0.490, abs=0.01) and hi == pytest.approx(0.943, abs=0.01)
    assert np.isnan(vc.wilson_ci(0, 0)[0])
    assert vc.wilson_ci(10, 10)[1] == 1.0


def test_precision_uses_only_the_stratified_sample():
    rows = []
    for i, (v, ins) in enumerate([("real", True), ("real", True), ("flicker", True),
                                  ("flicker", False), ("flicker", False)]):
        rows.append({"sample_id": f"S01.{i + 1}", "contact_type": "N2N", "verdict": v, "in_sample": ins,
                     "duration_sec": 0.5, "peak_score": 0.9, "duration_tertile": "short"})
    s = vc.compute_summary(pd.DataFrame(rows), vc.DEFAULT_CONFIG)
    assert s["n_reviewed"] == 3 and s["per_type"]["N2N"]["precision"] == pytest.approx(2 / 3)
    assert s["extras"]["n_reviewed"] == 2 and s["extras"]["per_type"]["N2N"]["precision"] == 0.0


def test_score_on_synthetic_review(tmp_path):
    verdicts = ["real", "real", "flicker", "wrong_type", "unsure", "tracking_error", "real", "flicker"]
    rows = []
    for i in range(len(verdicts)):
        rows.append({"sample_id": f"V{i + 1:03d}", "clip": f"clips/{i}.mp4", "thumbnail": "t.png",
                     "pair_label": "R1-R2", "contact_type": "N2N" if i < 5 else "SBS",
                     "contact_label": "x", "start_time": "00:00.0", "end_time": "00:01.0",
                     "duration_sec": 0.2 + 0.4 * i, "duration_frames": 5,
                     "duration_tertile": vc.TERTILES[i % 3], "peak_score": 0.5 + 0.05 * i})
    p = tmp_path / "review.xlsx"
    vc.write_review_xlsx(rows, p)
    df = pd.read_excel(p, sheet_name="review")
    df["verdict"] = verdicts
    df["correct_type"] = [None, None, None, "N2B", None, None, None, None]
    with pd.ExcelWriter(p) as w:
        df.to_excel(w, sheet_name="review", index=False)
    assert vc.main(["score", str(p)]) == 0
    s = json.loads((tmp_path / "validation_summary.json").read_text())
    n2n = s["per_type"]["N2N"]
    assert n2n["n_reviewed"] == 5 and n2n["n_real"] == 2 and n2n["n_decided"] == 4
    assert n2n["precision"] == 0.5
    assert s["wrong_type_confusion"] == {"N2N": {"N2B": 1}}
    assert (tmp_path / "validation_summary.md").exists()
    assert (tmp_path / "validation_precision.png").exists()


def _ev(pk, ct, a, b):
    return {"pair_key": pk, "contact_type": ct, "start_frame": a, "end_frame": b, "n_frames": b - a + 1,
            "duration_sec": (b - a + 1) / 30, "peak_score": 0.9, "duration_tertile": "short"}


def test_build_scenes_groups_overlapping_events_and_adds_extras():
    events = pd.DataFrame([_ev("0_1", "N2N", 10, 40), _ev("1_2", "SBS", 30, 60),
                           _ev("0_2", "N2AG", 45, 80), _ev("0_1", "N2N", 200, 230)])
    sample = events.iloc[[0, 1, 3]]            # the N2AG was not sampled
    scenes = vc.build_scenes(sample, events, pad=15, nframes=1000)
    assert [s["scene_id"] for s in scenes] == ["S01", "S02"]
    s1 = scenes[0]
    assert (s1["a"], s1["b"]) == (0, 95)       # grown to cover the extra N2AG (80 + 15)
    assert list(s1["events"]["contact_type"]) == ["N2N", "SBS", "N2AG"]
    assert list(s1["events"]["in_sample"]) == [True, True, False]
    assert not s1["events"]["partial"].any()
    assert len(scenes[1]["events"]) == 1 and (scenes[1]["a"], scenes[1]["b"]) == (185, 245)


def test_build_scenes_splits_a_long_chain_into_non_overlapping_segments():
    events = pd.DataFrame([_ev("0_1", "N2N", 10, 40), _ev("1_2", "SBS", 30, 60), _ev("0_2", "N2AG", 45, 80)])
    scenes = vc.build_scenes(events.iloc[[0, 1]], events, pad=15, nframes=1000, max_frames=76)
    # chain [0, 95] = 96 frames > 76 -> two segments of 48 frames, no repeated footage
    assert [(s["a"], s["b"]) for s in scenes] == [(0, 47), (48, 95)]
    s1, s2 = scenes
    assert list(s1["events"]["contact_type"]) == ["N2N", "SBS"]      # most of each lies in segment 1
    assert list(s1["events"]["partial"]) == [False, True]            # SBS continues into segment 2
    assert list(s2["events"]["contact_type"]) == ["N2AG"]
    assert bool(s2["events"].iloc[0]["partial"]) is True and not bool(s2["events"].iloc[0]["in_sample"])
    assert list(s2["visible"]["contact_type"]) == ["SBS", "N2AG"]   # SBS still highlighted there
    # each event has exactly one review row
    keys = [k for s in scenes for k in vc._event_keys(s["events"])]
    assert len(keys) == len(set(keys)) == 3


def test_build_scenes_splits_sampled_events_too():
    events = pd.DataFrame([_ev("0_1", "N2N", 10, 40), _ev("1_2", "SBS", 30, 60)])
    scenes = vc.build_scenes(events, events, pad=15, nframes=1000, max_frames=50)
    assert [(s["a"], s["b"]) for s in scenes] == [(0, 37), (38, 75)]
    assert [list(s["events"]["contact_type"]) for s in scenes] == [["N2N"], ["SBS"]]
    assert all(s["events"]["partial"].all() for s in scenes)


# ------------------------------------------------------------------ NC windows (recall proxy)
FPS = 30.0


def _raw_events_table():
    """Events table as the post-processing writes it: NC rows fill the gaps between contacts.

    Pair 1_2 is deliberately inconsistent: one NC row covers a contact event (N2AG 300-360),
    which the sampler must still never hand out.
    """
    rows = [
        ("0_1", "NC", 0, 99), ("0_1", "N2N", 100, 159), ("0_1", "NC", 160, 399),
        ("0_1", "SBS", 400, 430), ("0_1", "NC", 431, 445), ("0_1", "FOL", 446, 480),
        ("0_1", "NC", 481, 799),
        ("0_2", "NC", 0, 799),
        ("1_2", "NC", 0, 799), ("1_2", "N2AG", 300, 360),
    ]
    return pd.DataFrame([{"pair_key": pk, "contact_type": ct, "start_frame": a, "end_frame": b,
                          "duration_sec": (b - a + 1) / FPS, "peak_score": 0.9 if ct != "NC" else np.nan}
                         for pk, ct, a, b in rows])


def _overlaps(w, contacts):
    same = contacts[contacts["pair_key"] == w["pair_key"]]
    return bool(((same["start_frame"] <= w["end_frame"]) & (same["end_frame"] >= w["start_frame"])).any())


def test_nc_spans_are_the_nc_rows_and_load_events_drops_them(tmp_path):
    raw = _raw_events_table()
    spans = vc.nc_spans(raw)
    assert len(spans) == 6 and set(spans.columns) == {"pair_key", "start_frame", "end_frame"}
    cdir = tmp_path / "contacts"
    cdir.mkdir()
    raw.to_csv(cdir / "contacts_real_events.csv", index=False)
    events, src = vc.load_events(tmp_path)
    assert src == "events" and set(events["contact_type"]) == {"N2N", "SBS", "FOL", "N2AG"}
    assert len(vc.load_nc_spans(tmp_path)) == 6


@pytest.mark.parametrize("seed", range(12))
def test_nc_windows_never_overlap_a_contact_event_of_the_same_pair(seed):
    raw = _raw_events_table()
    contacts, nc = vc._normalise_events(raw), vc.nc_spans(raw)
    wins = vc.sample_nc_windows(nc, contacts, None, FPS, 40, 3.0, 2.0, seed)
    assert len(wins) > 5 and (wins["contact_type"] == "NC").all()
    assert wins["nc_near"].isna().all()                          # no per-frame table: uniform
    for _, w in wins.iterrows():
        assert not _overlaps(w, contacts), dict(w)
        assert 60 <= w["n_frames"] <= 90                         # >= 2 * pad_sec, <= nc_window_sec
        assert not (w["pair_key"] == "0_1" and 431 <= w["start_frame"] <= 445)   # 15-frame span: too short
    for _, g in wins.groupby("pair_key"):                        # windows of a pair do not overlap
        g = g.sort_values("start_frame")
        assert (g["start_frame"].to_numpy()[1:] > g["end_frame"].to_numpy()[:-1]).all()


def test_nc_windows_are_seeded_and_capped():
    raw = _raw_events_table()
    contacts, nc = vc._normalise_events(raw), vc.nc_spans(raw)
    a = vc.sample_nc_windows(nc, contacts, None, FPS, 5, 3.0, 2.0, 1)
    b = vc.sample_nc_windows(nc, contacts, None, FPS, 5, 3.0, 2.0, 1)
    c = vc.sample_nc_windows(nc, contacts, None, FPS, 5, 3.0, 2.0, 2)
    assert len(a) == 5 and a.equals(b) and not a.equals(c)
    assert len(vc.sample_nc_windows(nc, contacts, None, FPS, 1000, 3.0, 2.0, 1)) < 1000   # only what exists
    assert vc.sample_nc_windows(nc, contacts, None, FPS, 0, 3.0, 2.0, 1).empty


def test_nc_windows_prefer_pairs_that_come_close():
    raw = _raw_events_table()
    contacts, nc = vc._normalise_events(raw), vc.nc_spans(raw)
    pf = pd.DataFrame([{"pair_key": pk, "frame_idx": f,
                        "centroid_dist_bl": 0.5 if (pk == "0_2" and 600 <= f <= 620) else 2.5}
                       for pk in ("0_1", "0_2", "1_2") for f in range(800)])
    pf_pairs = vc.index_per_frame(pf)
    for seed in range(5):
        one = vc.sample_nc_windows(nc, contacts, pf_pairs, FPS, 1, 3.0, 2.0, seed, near_bl=1.0)
        assert len(one) == 1 and bool(one.iloc[0]["nc_near"])
        assert one.iloc[0]["pair_key"] == "0_2" and one.iloc[0]["start_frame"] <= 600 <= one.iloc[0]["end_frame"]
        assert one.iloc[0]["min_centroid_dist_bl"] == pytest.approx(0.5)
    many = vc.sample_nc_windows(nc, contacts, pf_pairs, FPS, 6, 3.0, 2.0, 0)
    assert len(many) == 6 and int(many["nc_near"].sum()) == 1     # the single near window, then filled up
    assert not any(_overlaps(w, contacts) for _, w in many.iterrows())


def _nc_review_rows():
    """3 contact rows (N2N: real, real, flicker) and 4 NC windows (2 correct, 1 missed, 1 unsure)."""
    rows = []
    for i, (ct, v) in enumerate([("N2N", "real"), ("N2N", "real"), ("N2N", "flicker"),
                                 ("NC", "nc_correct"), ("NC", "nc_correct"), ("NC", "missed_contact"),
                                 ("NC", "unsure")]):
        nc = ct == "NC"
        rows.append({"sample_id": f"N{i - 2:02d}.1" if nc else f"S01.{i + 1}", "clip": f"clips/{i}.mp4",
                     "thumbnail": "t.png", "in_sample": True, "pair_label": "R1-R2", "contact_type": ct,
                     "contact_label": "x", "start_time": "00:00.0", "end_time": "00:01.0",
                     "duration_sec": 3.0 if nc else 0.5, "duration_frames": 15,
                     "duration_tertile": "" if nc else "short", "peak_score": None if nc else 0.9,
                     "nc_near": True if nc else None, "_verdict": v})
    return rows


def _fill(path, rows, verdicts, correct_type=None):
    vc.write_review_xlsx([{k: v for k, v in r.items() if k != "_verdict"} for r in rows], path)
    df = pd.read_excel(path, sheet_name="review")
    df["verdict"] = verdicts
    if correct_type is not None:
        df["correct_type"] = correct_type
    with pd.ExcelWriter(path) as w:
        df.to_excel(w, sheet_name="review", index=False)


def test_review_workbook_offers_the_nc_verdicts(tmp_path):
    from openpyxl import load_workbook
    p = tmp_path / "review.xlsx"
    vc.write_review_xlsx([{k: v for k, v in r.items() if k != "_verdict"} for r in _nc_review_rows()], p)
    wb = load_workbook(p)
    lists = [dv.formula1 for dv in wb["review"].data_validations.dataValidation]
    union = [f for f in lists if "real" in f]
    assert len(union) == 1
    for v in vc.ALL_VERDICTS:
        assert v in union[0]
    assert {"nc_correct", "missed_contact", "unsure"} <= set(vc.ALL_VERDICTS)
    assert "missed_contact" in "\n".join(str(c.value) for c in wb["Instructions"]["A"])


def test_nc_windows_give_a_recall_proxy_and_leave_precision_alone(tmp_path):
    rows = _nc_review_rows()
    p = tmp_path / "review.xlsx"
    _fill(p, rows, [r["_verdict"] for r in rows], [None, None, None, None, None, "N2B", None])
    assert vc.main(["score", str(p)]) == 0
    s = json.loads((tmp_path / "validation_summary.json").read_text())
    rp = s["recall_proxy"]
    assert (rp["n_nc_windows"], rp["n_missed"], rp["n_unsure"], rp["n_decided"]) == (4, 1, 1, 3)
    assert rp["miss_rate"] == pytest.approx(1 / 3)
    # Wilson 95% interval of 1/3, by hand: centre (p + z2/2n) / (1 + z2/n) = 0.4269, half-width 0.3654
    assert rp["ci_low"] == pytest.approx(0.0615, abs=1e-4) and rp["ci_high"] == pytest.approx(0.7923, abs=1e-4)
    assert rp["missed_by_type"] == {"N2B": 1} and rp["n_near_windows"] == 4
    # contact figures see only the three contact rows
    assert s["n_reviewed"] == 3 and s["per_type"]["N2N"]["precision"] == pytest.approx(2 / 3)
    assert "NC" not in s["per_type"] and s["extras"]["n_reviewed"] == 0
    md = (tmp_path / "validation_summary.md").read_text()
    assert "Recall proxy" in md and "33%" in md and "N2B x1" in md


def test_score_with_only_nc_rows(tmp_path):
    rows = [r for r in _nc_review_rows() if r["contact_type"] == "NC"]
    p = tmp_path / "review.xlsx"
    _fill(p, rows, [r["_verdict"] for r in rows])
    assert vc.main(["score", str(p)]) == 0
    s = json.loads((tmp_path / "validation_summary.json").read_text())
    assert s["n_reviewed"] == 0 and s["recall_proxy"]["miss_rate"] == pytest.approx(1 / 3)
    assert s["recall_proxy"]["missed_by_type"] == {"UNSPECIFIED": 1}


def test_score_without_nc_rows_has_no_recall_proxy():
    rows = [{"sample_id": "S01.1", "contact_type": "N2N", "verdict": "real", "in_sample": True,
             "duration_sec": 0.5, "peak_score": 0.9, "duration_tertile": "short"}]
    s = vc.compute_summary(pd.DataFrame(rows), vc.DEFAULT_CONFIG)
    assert s["recall_proxy"] is None and "Recall proxy" not in vc.render_md(s)


# ------------------------------------------------------------------ inter-rater agreement
def _agreement_rows():
    """(sample_id, contact_type, rater A, rater B): 8 contact rows with 6 binary agreements, one
    row rater B called unsure, and two NC rows."""
    return [
        ("S01.1", "N2N", "real", "real"), ("S01.2", "N2N", "real", "real"),
        ("S01.3", "N2N", "real", "flicker"), ("S01.4", "N2N", "flicker", "wrong_type"),
        ("S02.1", "SBS", "real", "real"), ("S02.2", "SBS", "real", "real"),
        ("S02.3", "SBS", "flicker", "real"), ("S02.4", "SBS", "tracking_error", "flicker"),
        ("S03.1", "FOL", "real", "unsure"),
        ("N01.1", "NC", "nc_correct", "nc_correct"), ("N02.1", "NC", "nc_correct", "nc_correct"),
    ]


def _write_review(path, rows, col):
    """review.xlsx filled with column `col` of `rows` (2 = rater A, 3 = rater B)."""
    path.parent.mkdir(parents=True, exist_ok=True)
    recs = [{"sample_id": r[0], "clip": "c.mp4", "pair_label": "R1-R2", "contact_type": r[1],
             "duration_sec": 0.5, "peak_score": 0.9, "duration_tertile": "short", "in_sample": True} for r in rows]
    vc.write_review_xlsx(recs, path)
    df = pd.read_excel(path, sheet_name="review")
    df["verdict"] = [r[col] for r in rows]
    df["reviewer"] = "AA" if col == 2 else "BB"
    with pd.ExcelWriter(path) as w:
        df.to_excel(w, sheet_name="review", index=False)


def test_cohen_kappa_matches_the_hand_computation():
    from fractions import Fraction
    a = [True, True, True, True, True, False, False, False]      # 5 yes, 3 no
    b = [True, True, True, True, False, True, False, False]      # 5 yes, 3 no; 6 of 8 agree
    st = vc.cohen_kappa(a, b)
    assert st["table"] == [[4, 1], [1, 2]] and st["n_agree"] == 6 and st["raw_agreement"] == 0.75
    p_o, p_a, p_b = Fraction(6, 8), Fraction(5, 8), Fraction(5, 8)
    p_e = p_a * p_b + (1 - p_a) * (1 - p_b)                      # 34/64
    expected = (p_o - p_e) / (1 - p_e)
    assert expected == Fraction(7, 15)
    assert st["kappa"] == pytest.approx(float(expected)) and st["p_expected"] == pytest.approx(float(p_e))
    assert vc.cohen_kappa([True, True], [True, True])["kappa"] is None        # one class only: undefined
    assert vc.cohen_kappa([], [])["kappa"] is None
    assert vc.cohen_kappa([True, False, True], [True, False, True])["kappa"] == 1.0
    assert vc.cohen_kappa([True, False], [False, True])["kappa"] == -1.0


def test_compare_two_review_workbooks(tmp_path):
    rows = _agreement_rows()
    pa, pb = tmp_path / "a" / "review.xlsx", tmp_path / "b" / "review.xlsx"
    _write_review(pa, rows, 2)
    _write_review(pb, rows[:-3] + rows[-2:], 3)         # rater B never saw S03.1
    assert vc.main(["score", "--compare", str(pa), str(pb)]) == 0
    r = json.loads((tmp_path / "a" / "interrater_summary.json").read_text())
    assert (tmp_path / "a" / "interrater_summary.md").exists()
    assert not (tmp_path / "b" / "interrater_summary.json").exists()           # written next to review A
    assert (r["n_matched"], r["n_only_a"], r["n_only_b"], r["n_unsure_excluded"]) == (10, 1, 0, 0)
    assert (r["reviewer_a"], r["reviewer_b"]) == ("AA", "BB")
    ov = r["overall"]                                    # contact rows only (NC rows are their own group)
    assert ov["n"] == 8 and ov["n_agree"] == 6 and ov["raw_agreement"] == pytest.approx(0.75)
    assert ov["table"] == [[4, 1], [1, 2]]
    assert ov["kappa"] == pytest.approx(7 / 15)         # (0.75 - 34/64) / (1 - 34/64), by hand
    assert r["per_type"]["N2N"]["kappa"] == pytest.approx(0.5)     # p_o 3/4, p_e 1/2
    assert r["per_type"]["SBS"]["kappa"] == pytest.approx(0.5)
    assert r["per_type"]["N2N"]["raw_agreement_full_verdict"] == pytest.approx(0.5)   # flicker vs wrong_type differ
    assert r["per_type"]["NC"]["raw_agreement"] == 1.0 and r["per_type"]["NC"]["kappa"] is None
    assert sorted(d["sample_id"] for d in r["disagreements"]) == ["S01.3", "S02.3"]
    md = (tmp_path / "a" / "interrater_summary.md").read_text()
    assert "0.467" in md and "undefined" in md and "S02.3" in md


def test_compare_excludes_unsure_rows(tmp_path):
    rows = _agreement_rows()
    pa, pb = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
    _write_review(pa, rows, 2)
    _write_review(pb, rows, 3)
    r = vc.compare_reviews(pd.read_excel(pa, sheet_name="review"), pd.read_excel(pb, sheet_name="review"))
    assert r["n_matched"] == 11 and r["n_unsure_excluded"] == 1 and r["overall"]["n"] == 8
    assert "FOL" not in r["per_type"]


def test_score_needs_exactly_one_of_review_and_compare(tmp_path):
    with pytest.raises(SystemExit):
        vc.main(["score"])
    with pytest.raises(SystemExit):
        vc.main(["score", str(tmp_path / "a.xlsx"), "--compare", "x.xlsx", "y.xlsx"])

# ------------------------------------------------------------------ investigator, boundaries, bouts, identity
def _ev_series(sid, pk, ct, a, b, inv=np.nan):
    """An event row as `sample` hands it to _event_row (investigator_slot is the pair slot, 0-based)."""
    return pd.Series({"sample_id": sid, "pair_key": pk, "contact_type": ct, "start_frame": a, "end_frame": b,
                      "n_frames": b - a + 1, "duration_sec": (b - a + 1) / FPS, "peak_score": 0.9,
                      "duration_tertile": "short", "investigator_slot": inv})


def test_review_prefill_names_the_investigator_and_marks_the_rest_na():
    n2ag = vc.review_prefill(_ev_series("S01.1", "1_2", "N2AG", 10, 40, inv=2))
    assert n2ag == {"investigator_label": "R3", "investigator_ok": None, "boundaries_ok": None}
    assert vc.review_prefill(_ev_series("S01.2", "1_2", "FOL", 10, 40, inv=1.0))["investigator_label"] == "R2"
    n2n = vc.review_prefill(_ev_series("S01.3", "0_1", "N2N", 10, 40))
    assert n2n == {"investigator_label": None, "investigator_ok": "na", "boundaries_ok": None}
    assert vc.review_prefill(_ev_series("S01.4", "0_1", "SBS", 10, 40, inv=0))["investigator_ok"] == "na"
    unknown = vc.review_prefill(_ev_series("S01.5", "0_1", "N2B", 10, 40, inv=np.nan))
    assert unknown["investigator_label"] is None and unknown["investigator_ok"] == "na"
    stray = vc.review_prefill(_ev_series("S01.6", "0_1", "N2B", 10, 40, inv=4))      # not one of the pair
    assert stray["investigator_label"] is None and stray["investigator_ok"] == "na"
    nc = vc.review_prefill(_ev_series("N01.1", "0_1", "NC", 10, 40))
    assert nc == {"investigator_label": None, "investigator_ok": "na", "boundaries_ok": "na"}
    assert vc.investigator_arrow(_ev_series("x", "1_2", "N2AG", 1, 2, inv=2)) == "R3->R2"
    assert vc.investigator_arrow(_ev_series("x", "1_2", "N2N", 1, 2)) == ""


def test_review_workbook_has_the_investigator_and_boundary_columns(tmp_path):
    from openpyxl import load_workbook
    evs = [_ev_series("S01.1", "1_2", "N2AG", 10, 40, inv=2), _ev_series("S01.2", "0_1", "N2N", 50, 80),
           _ev_series("N01.1", "0_1", "NC", 100, 189)]
    rows = [{**vc._event_row(e, None, FPS), "clip": "clips/S01.mp4", "thumbnail": "t.png", "in_sample": True}
            for e in evs]
    p = tmp_path / "review.xlsx"
    vc.write_review_xlsx(rows, p, speed=0.5)
    wb = load_workbook(p)
    ws = wb["review"]
    head = [c.value for c in ws[1]]
    i = head.index("correct_type")
    assert head[i:i + 5] == ["correct_type", "investigator_label", "investigator_ok", "boundaries_ok", "reviewer"]
    assert head[-1] == "notes" and {"start_frame", "end_frame"} <= set(head)
    got = {r[head.index("sample_id")].value: {c: r[head.index(c)].value for c in
                                              ("investigator_label", "investigator_ok", "boundaries_ok")}
           for r in ws.iter_rows(min_row=2)}
    assert got["S01.1"] == {"investigator_label": "R3", "investigator_ok": None, "boundaries_ok": None}
    assert got["S01.2"] == {"investigator_label": None, "investigator_ok": "na", "boundaries_ok": None}
    assert got["N01.1"] == {"investigator_label": None, "investigator_ok": "na", "boundaries_ok": "na"}
    lists = [dv.formula1 for dv in ws.data_validations.dataValidation]
    assert lists.count('"yes,no,na"') == 2
    text = "\n".join(str(c.value) for c in wb["Instructions"]["A"])
    assert "investigator_ok" in text and "boundaries_ok" in text and "0.2 s" in text
    assert "recording times" in text and "0.5× speed" in text and "scrub the player to clip_from" in text
    back = pd.read_excel(p, sheet_name="review")           # what `score` reads: 'na' must survive pandas
    assert list(back.loc[back["sample_id"] == "S01.2", "investigator_ok"]) == ["na"]


def _accuracy_rows():
    """Real events with investigator / boundary answers, plus rows that must not count."""
    def r(sid, ct, v, inv, bnd, in_sample=True):
        return {"sample_id": sid, "contact_type": ct, "verdict": v, "investigator_ok": inv, "boundaries_ok": bnd,
                "in_sample": in_sample, "duration_sec": 0.5, "peak_score": 0.9, "duration_tertile": "short"}
    return [r("S01.1", "N2AG", "real", "yes", "yes"), r("S01.2", "N2AG", "real", "yes", "no"),
            r("S01.3", "N2AG", "real", "no", "yes"), r("S01.4", "N2B", "real", "yes", "yes"),
            r("S01.5", "FOL", "real", "yes", "yes"), r("S01.6", "FOL", "real", "no", "no"),
            r("S01.7", "N2N", "real", "na", "yes"), r("S01.8", "SBS", "real", "na", "no"),
            r("S01.9", "N2AG", "flicker", "no", "no"),              # not real: ignored
            r("S01.10", "N2AG", "real", None, None),                # not answered: ignored
            r("S01.11", "N2AG", "real", "no", "no", in_sample=False)]  # extra: stays out, like precision


def test_investigator_and_boundary_accuracy():
    s = vc.compute_summary(pd.DataFrame(_accuracy_rows()), vc.DEFAULT_CONFIG)
    inv, bnd = s["investigator_accuracy"], s["boundary_accuracy"]
    assert (inv["overall"]["n_scored"], inv["overall"]["n_yes"], inv["overall"]["n_no"]) == (6, 4, 2)
    assert inv["overall"]["accuracy"] == pytest.approx(4 / 6)
    lo, hi = vc.wilson_ci(4, 6)
    assert inv["overall"]["ci_low"] == pytest.approx(lo) and inv["overall"]["ci_high"] == pytest.approx(hi)
    assert set(inv["per_type"]) == {"N2AG", "N2B", "FOL"}              # N2N / SBS have no investigator
    assert inv["per_type"]["N2AG"]["accuracy"] == pytest.approx(2 / 3) and inv["per_type"]["N2B"]["accuracy"] == 1.0
    assert inv["per_type"]["FOL"]["accuracy"] == 0.5
    assert (bnd["overall"]["n_scored"], bnd["overall"]["n_yes"]) == (8, 5)
    assert bnd["overall"]["accuracy"] == pytest.approx(5 / 8)
    assert set(bnd["per_type"]) == {"N2AG", "N2B", "FOL", "N2N", "SBS"}
    assert bnd["per_type"]["N2AG"]["accuracy"] == pytest.approx(2 / 3) and bnd["per_type"]["SBS"]["accuracy"] == 0.0
    md = vc.render_md(s)
    assert "## Investigator and event boundaries" in md and "| investigator | N2AG | 3 | 2 | 1 | 67% |" in md
    assert "| boundaries | all | 8 | 5 | 3 | 62% |" in md


def test_accuracy_is_none_without_answers():
    rows = [{k: v for k, v in r.items() if k not in ("investigator_ok", "boundaries_ok")}
            for r in _accuracy_rows()]
    s = vc.compute_summary(pd.DataFrame(rows), vc.DEFAULT_CONFIG)           # an older review file
    assert s["investigator_accuracy"] is None and s["boundary_accuracy"] is None
    assert "Investigator and event boundaries" not in vc.render_md(s)
    blank = pd.DataFrame([{**r, "investigator_ok": None, "boundaries_ok": None} for r in _accuracy_rows()])
    assert vc.compute_summary(blank, vc.DEFAULT_CONFIG)["boundary_accuracy"] is None


def test_bout_status_rules():
    def rows(*vc_pairs):
        return [pd.Series({"verdict": v, "correct_type": c}) for v, c in vc_pairs]
    assert vc._bout_status(rows(("real", None), ("real", None))) == "real"
    assert vc._bout_status(rows(("real", None), ("wrong_type", "N2B"))) == "real"        # sub-type confusion
    assert vc._bout_status(rows(("real", None), ("flicker", None))) == "false"
    assert vc._bout_status(rows(("real", None), ("tracking_error", None))) == "false"
    for ct in ("SBS", "FOL", "NC"):
        assert vc._bout_status(rows(("real", None), ("wrong_type", ct))) == "false"
    assert vc._bout_status(rows(("real", None), (None, None))) == "incomplete"
    assert vc._bout_status(rows(("real", None), ("unsure", None))) == "incomplete"
    assert vc._bout_status(rows(("real", None), ("wrong_type", None))) == "incomplete"
    assert vc._bout_status(rows((None, None), ("flicker", None))) == "false"             # a failure wins


def _bout_run(tmp_path, with_frames=True, with_csv=True, hpc_manifest=True):
    """A run folder with investigation_bouts.csv, a manifest naming a path that does not exist (as after a
    download from Bunya) and a filled review: one real bout, one false, one incomplete, one not sampled."""
    run = tmp_path / "run"
    (run / "contacts").mkdir(parents=True)
    (run / "validation").mkdir()
    if with_csv:
        pd.DataFrame([
            {"bout_id": 0, "pair_key": "0_1", "pair_label": "R1-R2", "start_frame": 10, "end_frame": 60},
            {"bout_id": 1, "pair_key": "0_1", "pair_label": "R1-R2", "start_frame": 100, "end_frame": 140},
            {"bout_id": 2, "pair_key": "0_1", "pair_label": "R1-R2", "start_frame": 200, "end_frame": 230},
            {"bout_id": 3, "pair_key": "0_2", "pair_label": "R1-R3", "start_frame": 10, "end_frame": 60},
        ]).to_csv(run / "contacts" / "investigation_bouts.csv", index=False)
    if hpc_manifest:
        (run / "validation" / "manifest.json").write_text(
            json.dumps({"run_dir": "/home/uqcvaras/computer-vision-tracking/outputs/runs/x", "fps": 30.0}))
    spec = [  # sample_id, type, pair, first frame, last frame, verdict, correct_type
        ("S01.1", "N2AG", "R1-R2", 10, 30, "real", None), ("S01.2", "N2B", "R1-R2", 31, 60, "wrong_type", "N2AG"),
        ("S02.1", "N2N", "R1-R2", 100, 120, "real", None), ("S02.2", "N2AG", "R1-R2", 121, 140, "flicker", None),
        ("S03.1", "N2AG", "R1-R2", 200, 215, "real", None), ("S03.2", "N2B", "R1-R2", 216, 230, None, None),
        ("S04.1", "SBS", "R1-R2", 61, 99, "real", None),            # not investigation: never mapped to a bout
        ("S05.1", "N2N", "R2-R3", 10, 60, "real", None),            # no bout of the pair 1_2: unmapped
        ("N01.1", "NC", "R1-R2", 300, 389, "nc_correct", None)]
    recs = []
    for sid, ct, pl, a, b, v, c in spec:
        rec = {"sample_id": sid, "clip": "clips/x.mp4", "pair_label": pl, "contact_type": ct, "in_sample": True,
               "start_time": vc.fmt_time(a / FPS), "end_time": vc.fmt_time((b + 1) / FPS),
               "duration_sec": (b - a + 1) / FPS, "duration_frames": b - a + 1, "duration_tertile": "short",
               "peak_score": 0.9}
        if with_frames:
            rec.update(start_frame=a, end_frame=b)
        recs.append(rec)
    p = run / "validation" / "review.xlsx"
    vc.write_review_xlsx(recs, p)
    df = pd.read_excel(p, sheet_name="review")
    df["verdict"] = [r[5] for r in spec]
    df["correct_type"] = [r[6] for r in spec]
    with pd.ExcelWriter(p) as w:
        df.to_excel(w, sheet_name="review", index=False)
    return run, p


@pytest.mark.parametrize("with_frames", [True, False])
def test_bout_precision_from_the_run_next_to_the_review(tmp_path, with_frames):
    run, p = _bout_run(tmp_path, with_frames=with_frames)       # False: an older file, frames from the times
    assert vc.main(["score", str(p)]) == 0
    s = json.loads((run / "validation" / "validation_summary.json").read_text())
    bp = s["bout_precision"]
    assert (bp["n_bouts_scored"], bp["n_real"], bp["n_false"], bp["n_incomplete"]) == (2, 1, 1, 1)
    assert bp["precision"] == pytest.approx(0.5)
    lo, hi = vc.wilson_ci(1, 2)
    assert bp["ci_low"] == pytest.approx(lo) and bp["ci_high"] == pytest.approx(hi)
    assert (bp["n_bouts_total"], bp["n_bouts_with_events"], bp["n_rows_unmapped"]) == (4, 3, 1)
    status = {b["start_frame"]: b["status"] for b in bp["bouts"]}
    assert status == {10: "real", 100: "false", 200: "incomplete"}
    md = (run / "validation" / "validation_summary.md").read_text()
    assert "## Investigation-bout precision" in md and "| 3 of 4 | 2 | 1 | 1 | 1 | 50% |" in md
    assert "S02.1/S02.2" in md                                   # the false bout is named


def test_bout_precision_with_run_override_and_missing_csv(tmp_path, caplog):
    run, p = _bout_run(tmp_path, with_csv=False, hpc_manifest=False)
    with caplog.at_level("WARNING"):
        assert vc.main(["score", str(p)]) == 0
    s = json.loads((run / "validation" / "validation_summary.json").read_text())
    assert s["bout_precision"] is None and "Investigation-bout" not in (run / "validation" / "validation_summary.md").read_text()
    assert any("investigation_bouts.csv" in m for m in caplog.messages)
    assert s["per_type"]["N2AG"]["n_reviewed"] == 3              # everything else scored as before
    other, _ = _bout_run(tmp_path / "elsewhere")                 # a run that has the csv; the review is elsewhere
    moved = tmp_path / "moved" / "review.xlsx"
    moved.parent.mkdir()
    moved.write_bytes(p.read_bytes())
    assert vc.main(["score", str(moved), "--run", str(other)]) == 0
    assert json.loads((moved.parent / "validation_summary.json").read_text())["bout_precision"]["n_real"] == 1


def test_resolve_run_prefers_an_existing_manifest_dir(tmp_path):
    review = tmp_path / "run" / "validation" / "review.xlsx"
    review.parent.mkdir(parents=True)
    assert vc.resolve_run(review)[0] == tmp_path / "run"
    (review.parent / "manifest.json").write_text(json.dumps({"run_dir": "/no/such/run", "fps": 25.0}))
    run, manifest = vc.resolve_run(review)
    assert run == tmp_path / "run" and manifest["fps"] == 25.0
    other = tmp_path / "other"
    other.mkdir()
    (review.parent / "manifest.json").write_text(json.dumps({"run_dir": str(other)}))
    assert vc.resolve_run(review)[0] == other
    assert vc.resolve_run(review, str(tmp_path / "run"))[0] == (tmp_path / "run").resolve()


def test_score_still_works_on_a_review_without_the_new_columns(tmp_path):
    """An older review.xlsx has neither the new columns nor start_frame / end_frame."""
    verdicts = ["real", "real", "flicker", "unsure"]
    rows = [{"sample_id": f"V{i + 1:03d}", "clip": "c.mp4", "pair_label": "R1-R2", "contact_type": "N2AG",
             "start_time": "00:00.0", "end_time": "00:01.0", "duration_sec": 0.5, "duration_frames": 15,
             "duration_tertile": "short", "peak_score": 0.8, "in_sample": True} for i in range(4)]
    run = tmp_path / "run"
    (run / "validation").mkdir(parents=True)
    p = run / "validation" / "review.xlsx"
    vc.write_review_xlsx(rows, p)
    df = pd.read_excel(p, sheet_name="review").drop(columns=["investigator_label", "investigator_ok",
                                                             "boundaries_ok", "start_frame", "end_frame"],
                                                    errors="ignore")
    df["verdict"] = verdicts
    with pd.ExcelWriter(p) as w:
        df.to_excel(w, sheet_name="review", index=False)
    assert "investigator_ok" not in pd.read_excel(p, sheet_name="review").columns
    assert vc.main(["score", str(p)]) == 0
    s = json.loads((run / "validation" / "validation_summary.json").read_text())
    assert s["per_type"]["N2AG"]["n_real"] == 2 and s["per_type"]["N2AG"]["precision"] == pytest.approx(2 / 3)
    assert s["investigator_accuracy"] is None and s["boundary_accuracy"] is None
    assert s["bout_precision"] is None and s["identity_check"] is None


def test_index_html_puts_the_clip_position_before_the_video_time(tmp_path):
    rows = [{"sample_id": "S01.11", "scene_id": "S01", "clip": "clips/S01.mp4", "contact_type": "N2AG",
             "pair_label": "R4-R6", "investigator_label": "R6", "start_time": "00:04.9", "end_time": "00:05.3",
             "clip_from": "00:09.8", "clip_to": "00:10.5", "duration_sec": 0.37, "duration_tertile": "short",
             "peak_score": 0.94, "in_sample": True},
            {"sample_id": "S01.12", "scene_id": "S01", "clip": "clips/S01.mp4", "contact_type": "SBS",
             "pair_label": "R1-R2", "start_time": "00:06.1", "end_time": "00:06.5", "clip_from": "00:12.1",
             "clip_to": "00:12.9", "duration_sec": 0.4, "duration_tertile": "short", "peak_score": 0.6,
             "in_sample": True}]
    p = tmp_path / "index.html"
    vc.write_index_html(rows, p, "t", speed=0.5)
    doc = p.read_text(encoding="utf-8")
    line = next(x for x in doc.split("<li>") if x.startswith("<b>S01.11"))
    assert line.index("clip ") < line.index("video ")
    assert "R4-R6 (R6 &rarr; R4)" in line
    assert "clip 00:09.8&ndash;00:10.5 (0.5× speed)" in line and "video 00:04.9&ndash;00:05.3" in line
    assert "0.37 s, short" in line and "peak 0.94" in line
    sbs = next(x for x in doc.split("<li>") if x.startswith("<b>S01.12"))
    assert "&rarr;" not in sbs                                           # symmetric type: no investigator
    assert ("Clips play at 0.5× speed: 'clip' times are what the player shows, "
            "'video' times are the original recording times.") in doc
    vc.write_index_html(rows, p, "t")                                    # normal speed: no speed claim
    doc1 = p.read_text(encoding="utf-8")
    assert "×" not in doc1 and "clip times" not in doc1 and "'clip' times are what the player shows" in doc1


def test_identity_checkpoints():
    assert vc.identity_checkpoints(3600, 30.0, 60, 40) == [0, 1800]
    assert vc.identity_checkpoints(3601, 30.0, 60, 40) == [0, 1800, 3600]
    assert vc.identity_checkpoints(10, 5.0, 0.2, 40) == list(range(10))
    long = vc.identity_checkpoints(30 * 60 * 120, 30.0, 60, 40)           # 2 h every minute = 120 > 40
    assert len(long) == 40 and long[0] == 0 and long[-1] == 30 * 60 * 119
    assert all(b > a for a, b in zip(long, long[1:]))
    assert len(vc.identity_checkpoints(30 * 60 * 120, 30.0, 60, 0)) == 120    # no cap
    assert vc.identity_checkpoints(0, 30.0, 60, 40) == []


def test_identity_config_merges_partial_sections():
    cfg = vc.load_config([])
    assert cfg["identity"] == {"every_sec": 60, "max_frames": 40}
    cfg = vc.load_config(["identity.every_sec=30"])
    assert cfg["identity"] == {"every_sec": 30, "max_frames": 40}


def _identity_run(tmp_path, n=10, fps=5):
    """Run folder with an overlay video (n gray 128x128 frames) and a results.xlsx PerFrame sheet of 3 animals;
    in frame 5 animal 3 (R3) has no centroid."""
    cv2 = pytest.importorskip("cv2")
    run = tmp_path / "run"
    (run / "overlays").mkdir(parents=True)
    (run / "contacts").mkdir()
    wr = cv2.VideoWriter(str(run / "overlays" / "cutie_composite_test.avi"), cv2.VideoWriter_fourcc(*"MJPG"),
                         fps, (128, 128))
    assert wr.isOpened()
    for _ in range(n):
        wr.write(np.full((128, 128, 3), 128, np.uint8))
    wr.release()
    cen = {0: (30, 30), 1: (98, 30), 2: (64, 80)}
    recs = []
    for f in range(n):
        for pk, (i, j) in (("0_1", (0, 1)), ("0_2", (0, 2)), ("1_2", (1, 2))):
            ci, cj = cen[i], cen[j]
            if f == 5 and 2 in (i, j):
                if i == 2:
                    ci = (-1, -1)
                if j == 2:
                    cj = (-1, -1)
            recs.append({"frame_idx": f, "pair_key": pk, "contact_type": "NC", "centroid_i_x": ci[0],
                         "centroid_i_y": ci[1], "centroid_j_x": cj[0], "centroid_j_y": cj[1]})
    with pd.ExcelWriter(run / "contacts" / "results.xlsx") as w:
        pd.DataFrame(recs).to_excel(w, sheet_name="PerFrame", index=False)
    return run, cen


def test_identity_writes_checkpoint_images_html_and_workbook(tmp_path):
    cv2 = pytest.importorskip("cv2")
    run, cen = _identity_run(tmp_path)
    assert vc.main(["identity", str(run), "--every", "1"]) == 0           # 5 fps, 10 frames: frames 0 and 5
    d = run / "validation" / "identity"
    pngs = sorted(p.name for p in d.glob("identity_*.png"))
    assert pngs == ["identity_0000.00s.png", "identity_0001.00s.png"]
    img = cv2.imread(str(d / pngs[0]))
    assert img.shape == (128, 128, 3)
    red, green, blue = (0, 0, 255), (0, 255, 0), (255, 0, 0)               # BGR overlay colours of R1..R3
    for slot, col in enumerate((red, green, blue)):
        x, y = cen[slot]
        assert tuple(int(c) for c in img[y, x]) == col                     # the dot sits on the centroid
    img5 = cv2.imread(str(d / pngs[1]))
    assert max(abs(int(c) - 128) for c in img5[80, 64]) <= 3               # R3 has no centroid in frame 5: skipped
    assert tuple(int(c) for c in img5[30, 30]) == red
    page = (d / "identity.html").read_text(encoding="utf-8")
    assert all(n in page for n in pngs) and "R1" in page and "R3 (not drawn)" in page
    from openpyxl import load_workbook
    wb = load_workbook(d / "identity.xlsx")
    ws = wb["identity"]
    assert [c.value for c in ws[1]] == ["time", "frame", "image", "animal", "identity_ok", "notes"]
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    assert len(rows) == 2 * 3                                              # checkpoints x animals
    assert [r[3] for r in rows[:3]] == ["R1", "R2", "R3"] and [r[0] for r in rows[::3]] == ["00:00.0", "00:01.0"]
    assert [r[1] for r in rows[::3]] == [0, 5] and all(r[4] is None for r in rows)
    assert "not drawn" in rows[5][5] and rows[2][5] is None
    assert any(dv.formula1 == '"yes,no,unsure"' for dv in ws.data_validations.dataValidation)
    assert "same individual as at time 0" in "\n".join(str(c.value) for c in wb["Instructions"]["A"])
    assert vc.main(["identity", str(run), "--every", "0.2", "--max", "3", "--out", str(tmp_path / "o")]) == 0
    assert len(list((tmp_path / "o").glob("*.png"))) == 3


def test_identity_falls_back_to_the_raw_video(tmp_path):
    run, _ = _identity_run(tmp_path)
    (run / "overlays" / "cutie_composite_test.avi").rename(tmp_path / "raw.avi")
    assert vc.main(["identity", str(run), "--every", "1"]) == 1             # no overlay, no usable video_path
    (run / "config_used.yaml").write_text(f"video_path: {(tmp_path / 'raw.avi').as_posix()}\n")
    assert vc.main(["identity", str(run), "--every", "1"]) == 0
    assert len(list((run / "validation" / "identity").glob("*.png"))) == 2


def _identity_table(verdicts):
    """(time, animal) -> verdict rows for 3 checkpoints x 2 animals."""
    times = ["00:00.0", "01:00.0", "02:00.0"]
    return pd.DataFrame([{"time": t, "frame": k, "image": f"i{k}.png", "animal": a, "identity_ok": verdicts.get((t, a)),
                          "notes": None} for k, t in enumerate(times) for a in ("R1", "R2")])


def test_identity_check_summary():
    assert vc.identity_check(_identity_table({})) is None                   # nothing filled in yet
    ok = {(t, a): "yes" for t in ("00:00.0", "01:00.0", "02:00.0") for a in ("R1", "R2")}
    ok[("01:00.0", "R2")] = "no"
    ok[("02:00.0", "R2")] = "no"
    ok[("02:00.0", "R1")] = "unsure"
    ic = vc.identity_check(_identity_table(ok))
    assert (ic["n_checkpoints"], ic["n_checkpoints_ok"], ic["n_animals"]) == (3, 1, 2)
    assert (ic["n_checked"], ic["n_ok"], ic["n_no"], ic["n_unsure"]) == (6, 3, 2, 1)
    assert ic["first_failure_time_per_animal"] == {"R2": "01:00.0"}
    partial = vc.identity_check(_identity_table({("00:00.0", "R1"): "yes", ("00:00.0", "R2"): "yes"}))
    assert (partial["n_checkpoints"], partial["n_checkpoints_ok"]) == (1, 1)    # unfilled rows are not counted


def test_score_reports_identity_persistence(tmp_path):
    run, p = _bout_run(tmp_path)
    d = run / "validation" / "identity"
    d.mkdir()
    tbl = _identity_table({(t, a): "yes" for t in ("00:00.0", "01:00.0", "02:00.0") for a in ("R1", "R2")})
    tbl.loc[(tbl["time"] == "02:00.0") & (tbl["animal"] == "R2"), "identity_ok"] = "no"
    with pd.ExcelWriter(d / "identity.xlsx") as w:
        tbl.to_excel(w, sheet_name="identity", index=False)
    assert vc.main(["score", str(p)]) == 0
    s = json.loads((run / "validation" / "validation_summary.json").read_text())
    ic = s["identity_check"]
    assert (ic["n_checkpoints"], ic["n_checkpoints_ok"], ic["n_no"]) == (3, 2, 1)
    assert ic["first_failure_time_per_animal"] == {"R2": "02:00.0"}
    md = (run / "validation" / "validation_summary.md").read_text()
    assert "Identity persistence: 2/3 checkpoints ok" in md and "R2 swapped or lost at 02:00.0" in md


def test_bouts_source_falls_back_to_diagnostics_xlsx(tmp_path):
    """--source bouts: contact_bouts.csv, else Bouts of diagnostics.xlsx, else Bouts of results.xlsx (old runs)."""
    def bouts(tag):
        return pd.DataFrame([{"pair_key": "0_1", "contact_type": "N2N", "start_frame": 10, "end_frame": 40,
                              "duration_sec": 1.0, "peak_score": 0.9, "tag": tag}])

    def book(path, **sheets):
        with pd.ExcelWriter(path) as w:
            for name, df in sheets.items():
                df.to_excel(w, sheet_name=name, index=False)

    cdir = tmp_path / "contacts"
    cdir.mkdir()
    book(cdir / "results.xlsx", Summary=pd.DataFrame({"a": [1]}))            # consolidated: no Bouts sheet
    with pytest.raises(FileNotFoundError):
        vc.load_events(tmp_path, source="bouts")
    book(cdir / "diagnostics.xlsx", Bouts=bouts("diagnostics"), Dynamics=pd.DataFrame({"a": [1]}))
    ev, src = vc.load_events(tmp_path, source="bouts")
    assert src == "bouts" and list(ev["tag"]) == ["diagnostics"] and list(ev["contact_type"]) == ["N2N"]
    bouts("csv").to_csv(cdir / "contact_bouts.csv", index=False)               # the csv still wins
    assert list(vc.load_events(tmp_path, source="bouts")[0]["tag"]) == ["csv"]
    old = tmp_path / "old"                                                    # a run from before the split
    (old / "contacts").mkdir(parents=True)
    book(old / "contacts" / "results.xlsx", Bouts=bouts("results"))
    assert list(vc.load_events(old, source="bouts")[0]["tag"]) == ["results"]
