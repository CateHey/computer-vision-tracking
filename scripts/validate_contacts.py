#!/usr/bin/env python
"""Validation phase for detected social contacts.

Subcommands
-----------
sample : stratified sample of events -> evidence clips, contact sheets,
         review.xlsx, index.html, manifest.json. By default events that overlap
         in time are grouped into one clip (a scene) with every pair highlighted;
         review.xlsx keeps one row per event and says where to look in the clip.
         With --nc-per-run N it also samples N windows from the spans where the pipeline
         reported NO contact (recall proxy: a reviewer says whether a contact was missed).
         Every row also asks whether the named investigator and the event boundaries are right.
score  : read a filled review.xlsx -> validation_summary.{md,json,png}
         (precision per type; recall_proxy when NC windows were reviewed; investigator and
         boundary accuracy; investigation-bout precision from <run>/contacts/investigation_bouts.csv;
         identity persistence when <run>/validation/identity/identity.xlsx is filled).
         `score --compare A.xlsx B.xlsx` -> interrater_summary.{md,json} (Cohen's kappa).
identity : checkpoint frames every N seconds with each animal's label drawn from the per-frame
         centroids -> <run>/validation/identity/ (PNGs, identity.html, identity.xlsx), to check that
         the labels stay on the same individual over a long session.

Usage
-----
    python scripts/validate_contacts.py sample <run_dir> [--per-type 15] [--nc-per-run 20] [key=value ...]
    python scripts/validate_contacts.py score <validation_dir>/review.xlsx [--run <run_dir>]
    python scripts/validate_contacts.py score --compare review_A.xlsx review_B.xlsx
    python scripts/validate_contacts.py identity <run_dir> [--every 60] [--max 40]
"""
from __future__ import annotations

import argparse
import html
import json
import logging
import math
import shutil
import subprocess
import sys
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

import numpy as np
import pandas as pd

_PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(_PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(_PROJECT_ROOT))

logger = logging.getLogger(__name__)

DEFAULT_CONFIG: Dict[str, Any] = {
    "per_type": 15, "pad_sec": 1.0, "speed": 0.5, "seed": 42, "crf": 23,
    "source": None, "types": None,
    "group_by": "scene", "max_scene_sec": 12.0,
    "nc_windows": 0, "nc_window_sec": 3.0, "nc_near_bl": 1.0,
    "duration_bins": [0.3, 0.5, 1.0, 2.0],
    "score_grid": [0.0, 0.5, 0.6, 0.7, 0.8, 0.9],
    "duration_grid": [0.0, 0.2, 0.3, 0.5, 1.0, 2.0],
    "identity": {"every_sec": 60, "max_frames": 40},
}
TYPE_NAMES = {
    "N2N": "Nose-to-Nose", "N2AG": "Nose-to-Anogenital", "N2B": "Nose-to-Body",
    "SBS": "Side-by-Side", "FOL": "Following",
}
VERDICTS = ["real", "flicker", "wrong_type", "tracking_error", "unsure"]
NC_VERDICTS = ["nc_correct", "missed_contact", "unsure"]  # verdicts that apply to NC-window rows
# One dropdown for every row of review.xlsx: contact rows use VERDICTS, NC rows use NC_VERDICTS.
ALL_VERDICTS = ["real", "flicker", "wrong_type", "tracking_error", "nc_correct", "missed_contact", "unsure"]
CORRECT_TYPES = ["N2N", "N2AG", "N2B", "SBS", "FOL", "NC"]
INVESTIGATOR_TYPES = ("N2AG", "N2B", "FOL")      # asymmetric types: one animal investigates / follows
INVESTIGATION_TYPES = ("N2N", "N2AG", "N2B")     # the types merged into investigation bouts
YES_NO_NA = ["yes", "no", "na"]                  # investigator_ok / boundaries_ok
IDENTITY_VERDICTS = ["yes", "no", "unsure"]      # identity.xlsx
QUALITY_FLAGS = ["stale_keypoints", "missing_keypoints", "single_detection",
                 "high_mask_overlap", "fol_used_centroid", "nose_off_mask"]
TERTILES = ["short", "medium", "long"]
NO_CONTACT = {"", "none", "nc", "nan", "non_contact", "null"}


# --------------------------------------------------------------------------- config
def _set_nested(d: dict, dotted_key: str, value: Any) -> None:
    """Set a nested dict value using dotted key notation."""
    keys = dotted_key.split(".")
    for k in keys[:-1]:
        d = d.setdefault(k, {})
    d[keys[-1]] = value


def _parse_value(s: str) -> Any:
    """Parse a CLI override value to its natural type."""
    if s.lower() in ("true", "yes"):
        return True
    if s.lower() in ("false", "no"):
        return False
    if s.lower() in ("null", "none"):
        return None
    for cast in (int, float):
        try:
            return cast(s)
        except ValueError:
            pass
    if s.startswith("["):
        try:
            return json.loads(s)
        except ValueError:
            pass
    return s


def load_config(overrides: List[str]) -> dict:
    """Load configs/validation.yaml (or defaults) and apply key=value overrides."""
    import copy
    cfg = copy.deepcopy(DEFAULT_CONFIG)
    path = _PROJECT_ROOT / "configs" / "validation.yaml"
    if path.exists():
        import yaml
        for k, v in (yaml.safe_load(path.read_text()) or {}).items():
            if isinstance(v, dict) and isinstance(cfg.get(k), dict):
                cfg[k].update(v)         # a partial section (identity.every_sec only) keeps the other defaults
            else:
                cfg[k] = v
    for ov in overrides:
        if "=" not in ov:
            logger.warning("Ignoring invalid override (no '='): %s", ov)
            continue
        k, v = ov.split("=", 1)
        _set_nested(cfg, k.strip(), _parse_value(v.strip()))
    return cfg


# --------------------------------------------------------------------------- helpers
def pair_label(pair_key: str) -> str:
    """'0_1' -> 'R1-R2'."""
    try:
        i, j = str(pair_key).split("_")
        return f"R{int(i) + 1}-R{int(j) + 1}"
    except ValueError:
        return str(pair_key)


def fmt_time(sec: float) -> str:
    """Seconds -> mm:ss.s."""
    m, s = divmod(float(sec), 60)
    return f"{int(m):02d}:{s:04.1f}"


def parse_time(value: Any) -> Optional[float]:
    """'mm:ss.s' (the review.xlsx / identity.xlsx time format, minutes may exceed 59) -> seconds; None if unreadable."""
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return None
    if hasattr(value, "hour") and hasattr(value, "minute"):          # a cell Excel turned into a time
        return value.hour * 3600 + value.minute * 60 + value.second + getattr(value, "microsecond", 0) / 1e6
    try:
        m, s = str(value).strip().split(":")
        return int(m) * 60 + float(s)
    except ValueError:
        return None


def slot_label(slot: Any) -> str:
    """'R<n>' of a 0-based animal slot ('' when the slot is missing or not a number)."""
    try:
        v = float(slot)
    except (TypeError, ValueError):
        return ""
    return "" if math.isnan(v) or v < 0 else f"R{int(v) + 1}"


def investigator_of(ev: Any) -> Optional[int]:
    """Slot of the investigating animal of an event row (investigator_slot), None for symmetric or
    unknown. The slot must be one of the two animals of the pair."""
    if str(ev.get("contact_type", "")).strip().upper() not in INVESTIGATOR_TYPES:
        return None
    label = slot_label(ev.get("investigator_slot"))
    slot = int(label[1:]) - 1 if label else None
    return slot if slot in _pair_slots(str(ev.get("pair_key", ""))) else None


def investigator_arrow(ev: Any, sep: str = "->") -> str:
    """'R2->R3' (investigator -> the other animal of the pair), '' when the event has none."""
    inv = investigator_of(ev)
    if inv is None:
        return ""
    si, sj = _pair_slots(str(ev["pair_key"]))
    return f"{slot_label(inv)}{sep}{slot_label(sj if inv == si else si)}"


def speed_note(speed: float) -> str:
    """'0.5x speed' for the clip playback speed, '' at normal speed."""
    return "" if abs(float(speed) - 1.0) < 1e-9 else f"{float(speed):g}× speed"


def wilson_ci(k: int, n: int, z: float = 1.96) -> Tuple[float, float]:
    """Wilson score interval for a binomial proportion. (nan, nan) if n == 0."""
    if n <= 0:
        return (float("nan"), float("nan"))
    p = k / n
    d = 1 + z * z / n
    c = (p + z * z / (2 * n)) / d
    h = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / d
    return (max(0.0, c - h), min(1.0, c + h))


# --------------------------------------------------------------------------- loading
def _read_xlsx_sheet(path: Path, sheet: str) -> Optional[pd.DataFrame]:
    try:
        return pd.read_excel(path, sheet_name=sheet)
    except Exception:  # missing file / sheet
        return None


def _find_xlsx(run: Path) -> Optional[Path]:
    for p in (run / "contacts" / "results.xlsx", run / "results.xlsx"):
        if p.exists():
            return p
    return None


def _load_events_table(run: Path, source: Optional[str] = None) -> Tuple[pd.DataFrame, str]:
    """Raw events (preferred, if they have pair_key) or bouts table, NC rows included."""
    xlsx = _find_xlsx(run)
    cdir = run / "contacts"
    cands: List[Tuple[str, Optional[pd.DataFrame]]] = []
    if source in (None, "events"):
        ev = None
        if (cdir / "contacts_real_events.csv").exists():
            ev = pd.read_csv(cdir / "contacts_real_events.csv")
        if (ev is None or "pair_key" not in ev.columns) and xlsx:
            ev = _read_xlsx_sheet(xlsx, "Events")
        cands.append(("events", ev))
    if source in (None, "bouts"):
        bt = pd.read_csv(cdir / "contact_bouts.csv") if (cdir / "contact_bouts.csv").exists() else None
        if bt is None and xlsx:
            # contact_bouts.csv, else the Bouts sheet of diagnostics.xlsx (next to results.xlsx), else the
            # Bouts sheet of results.xlsx itself (runs from before the sheet moved to diagnostics.xlsx)
            for book in (xlsx.parent / "diagnostics.xlsx", xlsx):
                bt = _read_xlsx_sheet(book, "Bouts") if book.exists() else None
                if bt is not None:
                    break
        cands.append(("bouts", bt))
    for name, df in cands:
        if df is not None and len(df) and "pair_key" in df.columns:
            return df, name
    raise FileNotFoundError(f"No events/bouts table with pair_key found under {run}")


def load_events(run: Path, source: Optional[str] = None) -> Tuple[pd.DataFrame, str]:
    """Load contact events (preferred, if they have pair_key) or bouts; NC rows are dropped.

    Returns (df, source_name).
    """
    df, name = _load_events_table(run, source)
    return _normalise_events(df), name


def nc_spans(df: pd.DataFrame) -> pd.DataFrame:
    """The no-contact (NC) rows of a raw events table: pair_key, start_frame, end_frame."""
    cols = ["pair_key", "start_frame", "end_frame"]
    ct = df["contact_type"].astype(str).str.strip().str.lower()
    nc = df[ct.isin(NO_CONTACT)].copy()
    if nc.empty:
        return pd.DataFrame({"pair_key": pd.Series(dtype=str), "start_frame": pd.Series(dtype=int),
                             "end_frame": pd.Series(dtype=int)})
    nc["pair_key"] = nc["pair_key"].astype(str)
    nc["start_frame"] = nc["start_frame"].astype(int)
    nc["end_frame"] = nc["end_frame"].astype(int)
    return nc[cols].reset_index(drop=True)


def load_nc_spans(run: Path, source: Optional[str] = None) -> pd.DataFrame:
    """NC spans of the run's events table (empty when the table has no NC rows)."""
    df, _ = _load_events_table(run, source)
    return nc_spans(df)


def _normalise_events(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df["contact_type"] = df["contact_type"].astype(str).str.strip().str.upper()
    df = df[~df["contact_type"].str.lower().isin(NO_CONTACT)].reset_index(drop=True)
    df["pair_key"] = df["pair_key"].astype(str)
    df["start_frame"] = df["start_frame"].astype(int)
    df["end_frame"] = df["end_frame"].astype(int)
    if "n_frames" not in df.columns:
        df["n_frames"] = df["end_frame"] - df["start_frame"] + 1
    if "duration_sec" not in df.columns:
        df["duration_sec"] = np.nan
    if "peak_score" not in df.columns:
        df["peak_score"] = np.nan
    return df


def load_per_frame(run: Path) -> Optional[pd.DataFrame]:
    """Per-frame pair table (csv, else PerFrame sheet); None if unavailable."""
    p = run / "contacts" / "contacts_per_frame.csv"
    if p.exists():
        return pd.read_csv(p)
    xlsx = _find_xlsx(run)
    return _read_xlsx_sheet(xlsx, "PerFrame") if xlsx else None


def find_fps(run: Path, video: Optional[Path]) -> float:
    """fps from session_summary.json, config_used.yaml, or the video."""
    try:
        s = json.loads((run / "contacts" / "session_summary.json").read_text())
        if s.get("metadata", {}).get("fps"):
            return float(s["metadata"]["fps"])
    except Exception:
        pass
    try:
        import yaml
        c = yaml.safe_load((run / "config_used.yaml").read_text()) or {}
        for k in ("fps", "video_fps"):
            if isinstance(c.get(k), (int, float)):
                return float(c[k])
        for sect in c.values():
            if isinstance(sect, dict) and isinstance(sect.get("fps"), (int, float)):
                return float(sect["fps"])
    except Exception:
        pass
    if video is not None:
        import cv2
        cap = cv2.VideoCapture(str(video))
        f = cap.get(cv2.CAP_PROP_FPS)
        cap.release()
        if f > 0:
            return float(f)
    return 30.0


def find_overlay(run: Path) -> Path:
    """Locate the main overlay video (not the per-animal unitary ones)."""
    ov = run / "overlays"
    for pat in ("cutie_composite_*.avi", "*composite*.avi", "*.avi", "*.mp4"):
        for p in sorted(ov.glob(pat)):
            if "unitary" not in p.name:
                return p
    raise FileNotFoundError(f"No overlay video in {ov}")


# --------------------------------------------------------------------------- sampling
def assign_tertiles(durations: pd.Series) -> pd.Series:
    """Label each event short/medium/long by duration rank (ties broken by order)."""
    n = len(durations)
    ranks = durations.rank(method="first").astype(int) - 1
    idx = np.minimum((ranks * 3) // max(n, 1), 2)
    return idx.map(lambda i: TERTILES[int(i)])


def stratified_sample(df: pd.DataFrame, n: int, seed: int) -> pd.DataFrame:
    """Sample up to n events evenly across duration tertiles (seeded).

    Returns rows with a 'duration_tertile' column. Takes all if len(df) <= n;
    when a tertile has too few events, its quota is redistributed.
    """
    df = df.copy()
    df["duration_tertile"] = assign_tertiles(df["duration_sec"]) if len(df) else []
    if len(df) <= n:
        return df.sort_values("start_frame")
    rng = np.random.default_rng(seed)
    groups = {t: df[df["duration_tertile"] == t] for t in TERTILES}
    quota = {t: n // 3 + (1 if i < n % 3 else 0) for i, t in enumerate(TERTILES)}
    take = {t: min(quota[t], len(groups[t])) for t in TERTILES}
    spare = n - sum(take.values())
    while spare > 0:
        grew = False
        for t in TERTILES:
            if spare > 0 and take[t] < len(groups[t]):
                take[t] += 1
                spare -= 1
                grew = True
        if not grew:
            break
    parts = []
    for t in TERTILES:
        g = groups[t]
        if take[t]:
            pick = rng.choice(len(g), size=take[t], replace=False)
            parts.append(g.iloc[np.sort(pick)])
    return pd.concat(parts).sort_values("start_frame")


def index_per_frame(pf: Optional[pd.DataFrame]) -> Dict[str, pd.DataFrame]:
    """Per-pair view of the per-frame table, indexed by frame_idx (empty dict when pf is None)."""
    out: Dict[str, pd.DataFrame] = {}
    if pf is None:
        return out
    for k, g in pf.groupby(pf["pair_key"].astype(str)):
        out[k] = g.drop_duplicates("frame_idx").set_index("frame_idx").sort_index()
    return out


def _subtract_intervals(a: int, b: int, cuts: Sequence[Tuple[int, int]]) -> List[Tuple[int, int]]:
    """[a, b] minus the inclusive intervals in `cuts` (empty pieces are dropped)."""
    pieces: List[Tuple[int, int]] = []
    cur = a
    for c0, c1 in sorted(cuts):
        if c1 < cur:
            continue
        if c0 > b:
            break
        if c0 > cur:
            pieces.append((cur, c0 - 1))
        cur = max(cur, c1 + 1)
        if cur > b:
            break
    if cur <= b:
        pieces.append((cur, b))
    return pieces


def nc_candidate_windows(nc: pd.DataFrame, contacts: pd.DataFrame, fps: float, window_sec: float,
                         min_span_sec: float, nframes: Optional[int] = None) -> pd.DataFrame:
    """Cut the NC spans into candidate windows: (pair_key, start_frame, end_frame).

    A span is first trimmed by every contact event of the same pair (safety net: the
    events table should already be consistent), and spans shorter than `min_span_sec`
    are dropped. A span is then cut into consecutive windows of `window_sec` (the
    remainder is dropped); a span shorter than one window becomes one shorter window.
    """
    win = max(1, int(round(window_sec * fps)))
    min_len = max(1, int(math.ceil(min_span_sec * fps - 1e-9)))
    cuts: Dict[str, List[Tuple[int, int]]] = {}
    if len(contacts):
        for k, g in contacts.groupby(contacts["pair_key"].astype(str)):
            cuts[k] = list(zip(g["start_frame"].astype(int), g["end_frame"].astype(int)))
    out: List[Tuple[str, int, int]] = []
    for _, r in nc.iterrows():
        pk, a, b = str(r["pair_key"]), max(0, int(r["start_frame"])), int(r["end_frame"])
        if nframes is not None:
            b = min(b, nframes - 1)
        if b < a:
            continue
        for s0, s1 in _subtract_intervals(a, b, cuts.get(pk, [])):
            length = s1 - s0 + 1
            if length < min_len:
                continue
            if length < win:
                out.append((pk, s0, s1))
            else:
                out.extend((pk, s0 + i * win, s0 + (i + 1) * win - 1) for i in range(length // win))
    return pd.DataFrame(out, columns=["pair_key", "start_frame", "end_frame"])


def _min_centroid_bl(pp: Optional[pd.DataFrame], a: int, b: int) -> float:
    """Smallest finite centroid distance (body lengths) of a pair over frames a..b; nan if unknown."""
    if pp is None or "centroid_dist_bl" not in pp.columns:
        return float("nan")
    v = pd.to_numeric(pp.loc[a:b, "centroid_dist_bl"], errors="coerce").to_numpy(dtype=float)
    v = v[np.isfinite(v)]
    return float(v.min()) if len(v) else float("nan")


def sample_nc_windows(nc: pd.DataFrame, contacts: pd.DataFrame, pf_pairs: Optional[Dict[str, pd.DataFrame]],
                      fps: float, n: int, window_sec: float, min_span_sec: float, seed: int,
                      near_bl: float = 1.0, nframes: Optional[int] = None) -> pd.DataFrame:
    """Sample n no-contact windows (seeded), the recall-proxy side of the validation.

    Misses hide where the animals are close, so windows in which the pair's centroid
    distance drops below `near_bl` body lengths at some frame (column nc_near = True) are
    drawn first; if there are fewer than n of them the rest is filled from the other
    windows. Without a per-frame table (pf_pairs empty) every window is equally likely and
    nc_near is NaN. Returns the windows sorted by start_frame with the columns of an event
    table (contact_type "NC") plus nc_near and min_centroid_dist_bl.
    """
    cand = nc_candidate_windows(nc, contacts, fps, window_sec, min_span_sec, nframes)
    cols = ["pair_key", "contact_type", "start_frame", "end_frame", "n_frames", "duration_sec",
            "peak_score", "duration_tertile", "nc_near", "min_centroid_dist_bl"]
    if cand.empty or n <= 0:
        return pd.DataFrame(columns=cols)
    have_pf = bool(pf_pairs)
    mins = np.array([_min_centroid_bl((pf_pairs or {}).get(r.pair_key), r.start_frame, r.end_frame)
                     for r in cand.itertuples()]) if have_pf else np.full(len(cand), np.nan)
    near = np.where(np.isfinite(mins), mins < near_bl, False)
    rng = np.random.default_rng([int(seed), 7])
    pools = [np.flatnonzero(near), np.flatnonzero(~near)] if have_pf else [np.arange(len(cand))]
    pick: List[int] = []
    for pool in pools:
        k = min(n - len(pick), len(pool))
        if k > 0:
            pick.extend(int(i) for i in rng.choice(pool, size=k, replace=False))
    pick.sort()
    out = cand.iloc[pick].copy()
    out["contact_type"] = "NC"
    out["n_frames"] = out["end_frame"] - out["start_frame"] + 1
    out["duration_sec"] = out["n_frames"] / fps
    out["peak_score"] = np.nan
    out["duration_tertile"] = ""
    out["nc_near"] = near[pick] if have_pf else np.nan
    out["min_centroid_dist_bl"] = mins[pick]
    return out.sort_values(["start_frame", "pair_key"]).reset_index(drop=True)[cols]


# --------------------------------------------------------------------------- clips
def _score_col(ctype: str) -> str:
    return f"score_{ctype.lower()}"


def event_stats(ev: pd.Series, pf: Optional[pd.DataFrame]) -> Dict[str, Any]:
    """Per-event aggregates from the per-frame table."""
    out: Dict[str, Any] = {}
    if pf is None:
        return out
    sub = pf[(pf["pair_key"].astype(str) == ev["pair_key"])
             & (pf["frame_idx"] >= ev["start_frame"]) & (pf["frame_idx"] <= ev["end_frame"])]
    if sub.empty:
        return out
    sc = _score_col(ev["contact_type"])
    if sc in sub:
        out["mean_score"] = float(sub[sc].mean())
    if "secondary_type" in sub:
        sec = sub["secondary_type"].dropna().astype(str)
        sec = sec[~sec.str.lower().isin(NO_CONTACT)]
        if len(sec):
            out["secondary_type"] = sec.mode().iloc[0]
            if "secondary_score" in sub:
                out["secondary_score"] = float(sub.loc[sec.index, "secondary_score"].mean())
    for f in QUALITY_FLAGS:
        if f in sub:
            out[f"frac_{f}"] = float((sub[f].fillna(0) > 0).mean())
    for c in ("nose_nose_dist_bl", "mask_contact_bl"):
        if c in sub:
            out[f"mean_{c}"] = float(sub[c].mean())
    return out


# Per-animal BGR colours of the overlay video (configs: output.overlay_colors), R1..R6
DEFAULT_COLORS: List[Tuple[int, int, int]] = [
    (0, 0, 255), (0, 255, 0), (255, 0, 0), (255, 255, 0), (255, 0, 255), (0, 255, 255)]
COLOR_NAMES = {(0, 0, 255): "red", (0, 255, 0): "green", (255, 0, 0): "blue", (255, 255, 0): "cyan",
               (255, 0, 255): "magenta", (0, 255, 255): "yellow"}  # BGR tuples


def overlay_colors(run: Path) -> List[Tuple[int, int, int]]:
    """BGR colour per animal from config_used.yaml (output.overlay_colors), else the defaults."""
    try:
        import yaml
        cfg = yaml.safe_load((run / "config_used.yaml").read_text(encoding="utf-8")) or {}
        raw = (cfg.get("output") or {}).get("overlay_colors")
        if raw:
            return [tuple(int(v) for v in c) for c in raw]
    except Exception:  # noqa: BLE001  missing file or bad yaml: fall back
        pass
    return DEFAULT_COLORS


def _pair_slots(pair_key: str) -> Tuple[Optional[int], Optional[int]]:
    """'3_4' -> (3, 4); (None, None) if the key is not two slot indices."""
    try:
        i, j = str(pair_key).split("_")
        return int(i), int(j)
    except ValueError:
        return None, None


def _px(row: Optional[pd.Series], kx: str, ky: str) -> Optional[Tuple[int, int]]:
    """Pixel point from two per-frame columns; None when absent or -1 (unavailable)."""
    if row is None or kx not in row or ky not in row:
        return None
    x, y = row[kx], row[ky]
    if pd.isna(x) or pd.isna(y) or float(x) < 0 or float(y) < 0:
        return None
    return int(round(float(x))), int(round(float(y)))


def highlight_pair(frame: np.ndarray, ev: pd.Series, row: Optional[pd.Series],
                   colors: Optional[List[Tuple[int, int, int]]]) -> None:
    """Circle the two animals of the event's pair in their overlay colours (in place).

    Needs the centroid/nose pixel columns of contacts_per_frame.csv (runs from
    October 2026 on); older runs simply get no highlight. For N2N the two noses
    are joined by a line.
    """
    import cv2
    if row is None:
        return
    colors = colors or DEFAULT_COLORS
    si, sj = _pair_slots(ev["pair_key"])
    noses = []
    for slot, cen_k, nose_k, bl_k in (
            (si, ("centroid_i_x", "centroid_i_y"), ("nose_i_x", "nose_i_y"), "body_length_i_px"),
            (sj, ("centroid_j_x", "centroid_j_y"), ("nose_j_x", "nose_j_y"), "body_length_j_px")):
        col = colors[slot % len(colors)] if slot is not None else (255, 255, 255)
        cen, nose = _px(row, *cen_k), _px(row, *nose_k)
        bl = row.get(bl_k)
        radius = int(max(30, 0.75 * float(bl))) if bl is not None and pd.notna(bl) and float(bl) > 0 else 60
        if cen is not None:
            cv2.circle(frame, cen, radius, (0, 0, 0), 5, cv2.LINE_AA)
            cv2.circle(frame, cen, radius, col, 3, cv2.LINE_AA)
            if slot is not None:
                _put(frame, f"R{slot + 1}", (cen[0] - 16, cen[1] - radius - 10), 0.8, col, 2)
        if nose is not None:
            noses.append(nose)
    if len(noses) == 2 and ev["contact_type"] == "N2N":
        cv2.line(frame, noses[0], noses[1], (0, 0, 0), 4, cv2.LINE_AA)
        cv2.line(frame, noses[0], noses[1], (255, 255, 255), 2, cv2.LINE_AA)


def pair_colour_hint(pair_key: str, colors: Optional[List[Tuple[int, int, int]]]) -> str:
    """' (yellow / magenta)' for the banner when both colours have a name, else ''."""
    colors = colors or DEFAULT_COLORS
    names = []
    for slot in _pair_slots(pair_key):
        if slot is None:
            return ""
        names.append(COLOR_NAMES.get(tuple(colors[slot % len(colors)]), ""))
    return f" ({names[0]} / {names[1]})" if all(names) else ""


def _put(img, text, org, scale=0.55, color=(255, 255, 255), th=1):
    import cv2
    cv2.putText(img, text, org, cv2.FONT_HERSHEY_SIMPLEX, scale, (0, 0, 0), th + 2, cv2.LINE_AA)
    cv2.putText(img, text, org, cv2.FONT_HERSHEY_SIMPLEX, scale, color, th, cv2.LINE_AA)


def annotate(frame: np.ndarray, sid: str, ev: pd.Series, fps: float, fidx: int,
             row: Optional[pd.Series],
             colors: Optional[List[Tuple[int, int, int]]] = None) -> np.ndarray:
    """Burn banner, state border and live values into a frame."""
    import cv2
    h, w = frame.shape[:2]
    inside = ev["start_frame"] <= fidx <= ev["end_frame"]
    col = (0, 0, 220) if inside else (140, 140, 140)
    ctype = ev["contact_type"]
    ban = frame.copy()
    cv2.rectangle(ban, (0, 0), (w, 58), (0, 0, 0), -1)
    cv2.rectangle(ban, (0, h - 30), (w, h), (0, 0, 0), -1)
    frame = cv2.addWeighted(ban, 0.6, frame, 0.4, 0)
    highlight_pair(frame, ev, row, colors)
    pl = pair_label(ev["pair_key"]) + pair_colour_hint(ev["pair_key"], colors)
    _put(frame, f"{sid} | {ctype} {TYPE_NAMES.get(ctype, '')} | look at {pl}", (8, 22), 0.6)
    _put(frame, f"event {fmt_time(ev['start_frame'] / fps)}-{fmt_time((ev['end_frame'] + 1) / fps)}  "
                f"frames {ev['start_frame']}-{ev['end_frame']}  dur {ev['n_frames'] / fps:.2f}s  "
                f"[{'LABEL ACTIVE' if inside else 'context'}] f={fidx}", (8, 48), 0.5,
         (200, 200, 255) if inside else (200, 200, 200))
    if row is not None:
        parts = [f"raw={row.get('contact_type', '?')}"]
        sc = _score_col(ctype)
        if sc in row and pd.notna(row[sc]):
            parts.append(f"{sc[6:]}={row[sc]:.2f}")
        for c, lab in (("nose_nose_dist_bl", "nn"), ("centroid_dist_bl", "cd"), ("mask_contact_bl", "mc")):
            if c in row and pd.notna(row[c]):
                parts.append(f"{lab}={row[c]:.2f}BL")
        _put(frame, "  ".join(parts), (8, h - 9), 0.5)
    cv2.rectangle(frame, (0, 0), (w - 1, h - 1), col, 8 if inside else 4)
    return frame


class _Writer:
    """H.264 writer via ffmpeg pipe, cv2 mp4v fallback."""

    def __init__(self, path: Path, w: int, h: int, fps: float, crf: int):
        self.proc, self.cv = None, None
        if shutil.which("ffmpeg"):
            self.proc = subprocess.Popen(
                ["ffmpeg", "-y", "-loglevel", "error", "-f", "rawvideo", "-pix_fmt", "bgr24",
                 "-s", f"{w}x{h}", "-r", f"{fps}", "-i", "-", "-c:v", "libx264",
                 "-pix_fmt", "yuv420p", "-crf", str(crf), "-movflags", "+faststart", str(path)],
                stdin=subprocess.PIPE)
        else:
            import cv2
            logger.warning("ffmpeg not found; falling back to cv2 mp4v")
            self.cv = cv2.VideoWriter(str(path), cv2.VideoWriter_fourcc(*"mp4v"), fps, (w, h))

    def write(self, f: np.ndarray) -> None:
        if self.proc:
            self.proc.stdin.write(np.ascontiguousarray(f).tobytes())
        else:
            self.cv.write(f)

    def close(self) -> None:
        if self.proc:
            self.proc.stdin.close()
            self.proc.wait()
        else:
            self.cv.release()


def cut_clip(cap, sid: str, ev: pd.Series, pf_pair: Optional[pd.DataFrame], fps: float,
             pad: int, nframes: int, out: Path, sheet: Path, speed: float, crf: int,
             colors: Optional[List[Tuple[int, int, int]]] = None) -> int:
    """Write the annotated clip and 3-thumbnail contact sheet. Returns frames written."""
    import cv2
    a = max(0, int(ev["start_frame"]) - pad)
    b = min(nframes - 1, int(ev["end_frame"]) + pad)
    cap.set(cv2.CAP_PROP_POS_FRAMES, a)
    wr = None
    keep: Dict[int, np.ndarray] = {}
    want = {int(ev["start_frame"]), (int(ev["start_frame"]) + int(ev["end_frame"])) // 2, int(ev["end_frame"])}
    n = 0
    for fi in range(a, b + 1):
        ok, fr = cap.read()
        if not ok:
            break
        h, w = fr.shape[:2]
        w2, h2 = w - w % 2, h - h % 2
        if wr is None:
            wr = _Writer(out, w2, h2, fps * speed, crf)
        row = pf_pair.loc[fi] if pf_pair is not None and fi in pf_pair.index else None
        fr = annotate(fr, sid, ev, fps, fi, row, colors)[:h2, :w2]
        wr.write(fr)
        if fi in want:
            keep[fi] = fr
        n += 1
    if wr:
        wr.close()
    if keep:
        thumbs = [cv2.resize(keep[k], None, fx=0.5, fy=0.5) for k in sorted(keep)]
        cv2.imwrite(str(sheet), np.hstack(thumbs))
    return n


# --------------------------------------------------------------------------- scenes
def _event_keys(df: pd.DataFrame) -> List[Tuple[str, str, int]]:
    return list(zip(df["pair_key"].astype(str), df["contact_type"].astype(str), df["start_frame"].astype(int)))


def build_scenes(sample: pd.DataFrame, events: pd.DataFrame, pad: int, nframes: int,
                 max_frames: Optional[int] = None) -> List[Dict[str, Any]]:
    """Group events into scenes (one clip each), one review row per event.

    1. Sampled events whose padded windows overlap form a chain; every other event of
       the run that overlaps the chain is pulled in (in_sample=False) until nothing
       else overlaps, so a clip never shows a contact the workbook does not list.
    2. A chain longer than max_frames is cut into consecutive, non-overlapping
       segments of equal length (no footage is repeated).
    3. Each event is listed in exactly one segment, the one holding most of it;
       partial=True when it continues outside that segment. Every event that is
       visible in a segment is still highlighted there (key "visible").
    """
    chains: List[Dict[str, Any]] = []
    for idx, e in sample.sort_values("start_frame").iterrows():
        a = max(0, int(e["start_frame"]) - pad)
        b = min(nframes - 1, int(e["end_frame"]) + pad)
        if chains and a <= chains[-1]["b"]:
            chains[-1]["b"] = max(chains[-1]["b"], b)
            chains[-1]["idx"].append(idx)
        else:
            chains.append({"a": a, "b": b, "idx": [idx]})

    out: List[Dict[str, Any]] = []
    for ch in chains:
        rows = sample.loc[ch["idx"]].copy()
        rows["in_sample"] = True
        a, b = ch["a"], ch["b"]
        for _ in range(50):  # extras can pull in further extras; stop at the fixpoint
            inside = events[(events["start_frame"] <= b) & (events["end_frame"] >= a)]
            have = set(_event_keys(rows))
            extra = inside[[k not in have for k in _event_keys(inside)]].copy()
            if extra.empty:
                break
            extra["in_sample"] = False
            rows = pd.concat([rows, extra])
            a = max(0, min(a, int(extra["start_frame"].min()) - pad))
            b = min(nframes - 1, max(b, int(extra["end_frame"].max()) + pad))
        rows = rows.sort_values(["start_frame", "pair_key"]).reset_index(drop=True)

        length = b - a + 1
        n_seg = 1 if not max_frames or length <= max_frames else -(-length // max_frames)
        seg_len = -(-length // n_seg)
        segments = [(a + i * seg_len, min(b, a + (i + 1) * seg_len - 1)) for i in range(n_seg)]

        def overlap(r, seg):
            return max(0, min(int(r["end_frame"]), seg[1]) - max(int(r["start_frame"]), seg[0]) + 1)

        assigned = [max(range(n_seg), key=lambda i: (overlap(r, segments[i]), -i)) for _, r in rows.iterrows()]
        for i, (sa, sb) in enumerate(segments):
            mine = rows[[k == i for k in assigned]].copy()
            mine["partial"] = (mine["start_frame"] < sa) | (mine["end_frame"] > sb)
            visible = rows[(rows["start_frame"] <= sb) & (rows["end_frame"] >= sa)].copy()
            out.append({"a": sa, "b": sb, "events": mine.reset_index(drop=True),
                        "visible": visible.reset_index(drop=True)})
    for k, sc in enumerate(out, 1):
        sc["scene_id"] = f"S{k:02d}"
    return out


def annotate_scene(frame: np.ndarray, scene_id: str, events: pd.DataFrame, n_listed: int, fps: float, fidx: int,
                   rows_by_pair: Dict[str, Optional[pd.Series]],
                   colors: Optional[List[Tuple[int, int, int]]]) -> np.ndarray:
    """Burn the scene banner, every active pair's highlight and the state border.

    `events` are all events visible in the clip (highlighted while active); n_listed is
    how many of them have their review row in this clip.
    """
    import cv2
    h, w = frame.shape[:2]
    active = events[(events["start_frame"] <= fidx) & (events["end_frame"] >= fidx)]
    inside = len(active) > 0
    col = (0, 0, 220) if inside else (140, 140, 140)
    ban = frame.copy()
    cv2.rectangle(ban, (0, 0), (w, 84), (0, 0, 0), -1)
    cv2.rectangle(ban, (0, h - 30), (w, h), (0, 0, 0), -1)
    frame = cv2.addWeighted(ban, 0.6, frame, 0.4, 0)
    for _, ev in active.iterrows():
        highlight_pair(frame, ev, rows_by_pair.get(str(ev["pair_key"])), colors)
    _put(frame, f"{scene_id} | {n_listed} event(s) to review in this clip | t={fmt_time(fidx / fps)} f={fidx} "
                f"[{'LABEL ACTIVE' if inside else 'context'}]", (8, 22), 0.6)
    labels = [f"{ev['sample_id']} {pair_label(ev['pair_key'])} {ev['contact_type']}"
              f"{' ' + investigator_arrow(ev) if investigator_arrow(ev) else ''}"
              f"{pair_colour_hint(ev['pair_key'], colors)}" for _, ev in active.iterrows()]
    _put(frame, "  |  ".join(labels[:2]) if labels else "no contact active", (8, 48), 0.5,
         (200, 200, 255) if inside else (200, 200, 200))
    if len(labels) > 2:
        _put(frame, "  |  ".join(labels[2:4]) + ("  ..." if len(labels) > 4 else ""), (8, 72), 0.5, (200, 200, 255))
    parts = []
    for _, ev in active.head(2).iterrows():
        row = rows_by_pair.get(str(ev["pair_key"]))
        if row is None:
            continue
        sc = _score_col(ev["contact_type"])
        txt = f"{pair_label(ev['pair_key'])} raw={row.get('contact_type', '?')}"
        if sc in row and pd.notna(row[sc]):
            txt += f" {sc[6:]}={row[sc]:.2f}"
        if "centroid_dist_bl" in row and pd.notna(row["centroid_dist_bl"]):
            txt += f" cd={row['centroid_dist_bl']:.2f}BL"
        parts.append(txt)
    if parts:
        _put(frame, "   ".join(parts), (8, h - 9), 0.5)
    cv2.rectangle(frame, (0, 0), (w - 1, h - 1), col, 8 if inside else 4)
    return frame


def annotate_nc_window(frame: np.ndarray, scene_id: str, win: pd.Series, fps: float, fidx: int,
                       row: Optional[pd.Series],
                       colors: Optional[List[Tuple[int, int, int]]]) -> np.ndarray:
    """Banner for a no-contact window: names the pair to watch, circles both animals, orange border."""
    import cv2
    h, w = frame.shape[:2]
    ban = frame.copy()
    cv2.rectangle(ban, (0, 0), (w, 58), (0, 0, 0), -1)
    cv2.rectangle(ban, (0, h - 30), (w, h), (0, 0, 0), -1)
    frame = cv2.addWeighted(ban, 0.6, frame, 0.4, 0)
    highlight_pair(frame, win, row, colors)
    pl = pair_label(win["pair_key"]) + pair_colour_hint(win["pair_key"], colors)
    _put(frame, f"{win['sample_id']} | NC window: is there any contact between {pl}?", (8, 22), 0.6)
    _put(frame, f"window {fmt_time(win['start_frame'] / fps)}-{fmt_time((win['end_frame'] + 1) / fps)}  "
                f"frames {win['start_frame']}-{win['end_frame']}  t={fmt_time(fidx / fps)} f={fidx}  "
                f"[no label active]", (8, 48), 0.5, (200, 200, 200))
    if row is not None:
        parts = [f"raw={row.get('contact_type', '?')}"]
        for c, lab in (("nose_nose_dist_bl", "nn"), ("centroid_dist_bl", "cd"), ("mask_contact_bl", "mc")):
            if c in row and pd.notna(row[c]):
                parts.append(f"{lab}={row[c]:.2f}BL")
        _put(frame, "  ".join(parts), (8, h - 9), 0.5)
    cv2.rectangle(frame, (0, 0), (w - 1, h - 1), (0, 165, 255), 6)
    return frame


def cut_scene_clip(cap, scene: Dict[str, Any], visible: pd.DataFrame, pf_pairs: Dict[str, pd.DataFrame], fps: float,
                   out: Path, sheet: Path, speed: float, crf: int,
                   colors: Optional[List[Tuple[int, int, int]]]) -> int:
    """Write one annotated clip for a scene (scene["kind"] == "nc": a no-contact window) and a
    3-thumbnail contact sheet."""
    import cv2
    a, b = scene["a"], scene["b"]
    pairs = [str(k) for k in visible["pair_key"].unique()]
    cap.set(cv2.CAP_PROP_POS_FRAMES, a)
    wr, keep, n = None, {}, 0
    want = {a, (a + b) // 2, b}
    for fi in range(a, b + 1):
        ok, fr = cap.read()
        if not ok:
            break
        h, w = fr.shape[:2]
        w2, h2 = w - w % 2, h - h % 2
        if wr is None:
            wr = _Writer(out, w2, h2, fps * speed, crf)
        rows_by_pair = {k: (pf_pairs[k].loc[fi] if k in pf_pairs and fi in pf_pairs[k].index else None)
                        for k in pairs}
        if scene.get("kind") == "nc":
            win = scene["events"].iloc[0]
            fr = annotate_nc_window(fr, scene["scene_id"], win, fps, fi,
                                    rows_by_pair.get(str(win["pair_key"])), colors)[:h2, :w2]
        else:
            fr = annotate_scene(fr, scene["scene_id"], visible, len(scene["events"]), fps, fi,
                                rows_by_pair, colors)[:h2, :w2]
        wr.write(fr)
        if fi in want:
            keep[fi] = fr
        n += 1
    if wr:
        wr.close()
    if keep:
        thumbs = [cv2.resize(keep[k], None, fx=0.5, fy=0.5) for k in sorted(keep)]
        cv2.imwrite(str(sheet), np.hstack(thumbs))
    return n


# --------------------------------------------------------------------------- outputs
REVIEWER_COLS = ["verdict", "correct_type", "investigator_label", "investigator_ok", "boundaries_ok",
                 "reviewer", "notes"]
COLUMN_ORDER = ["sample_id", "scene_id", "clip", "thumbnail", "in_sample", "partial",
                "pair_label", "contact_type", "contact_label", "nc_near", "min_centroid_dist_bl",
                "start_time", "end_time", "start_frame", "end_frame",
                "duration_sec", "duration_frames", "duration_tertile",
                "clip_from", "clip_to", "peak_score", "mean_score", "secondary_type", "secondary_score"] \
    + [f"frac_{f}" for f in QUALITY_FLAGS] + ["mean_nose_nose_dist_bl", "mean_mask_contact_bl"]


def review_prefill(ev: Any) -> Dict[str, Any]:
    """Pipeline claims pre-filled in the reviewer columns of one row.

    investigator_label: the animal the pipeline names as investigator (N2AG / N2B / FOL events).
    investigator_ok: left empty for the reviewer when there is an investigator to judge, 'na'
    otherwise (N2N, SBS, NC windows, or an event whose investigator is unknown).
    boundaries_ok: left empty for events, 'na' for NC windows (they have no event boundaries).
    """
    if str(ev["contact_type"]).strip().upper() == "NC":
        return {"investigator_label": None, "investigator_ok": "na", "boundaries_ok": "na"}
    inv = investigator_of(ev)
    return {"investigator_label": slot_label(inv) if inv is not None else None,
            "investigator_ok": None if inv is not None else "na", "boundaries_ok": None}


def write_review_xlsx(rows: List[Dict[str, Any]], path: Path, speed: float = 1.0) -> None:
    """review.xlsx with dropdowns + an Instructions sheet (`speed`: playback speed of the clips)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    present = set().union(*(r.keys() for r in rows)) if rows else set()
    cols = [c for c in COLUMN_ORDER if c in present] + REVIEWER_COLS
    wb = Workbook()
    ws = wb.active
    ws.title = "review"
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
    for r in rows:
        ws.append([r.get(c) for c in cols])
        cell = ws.cell(row=ws.max_row, column=cols.index("clip") + 1)
        cell.hyperlink = r["clip"]
        cell.font = Font(color="0563C1", underline="single")
    ws.freeze_panes = "B2"
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = 38 if c in ("clip", "thumbnail", "notes") else (
            22 if c == "contact_label" else 14)
    last = max(len(rows) + 1, 2)
    for col, opts in (("verdict", ALL_VERDICTS), ("correct_type", CORRECT_TYPES),
                      ("investigator_ok", YES_NO_NA), ("boundaries_ok", YES_NO_NA)):
        dv = DataValidation(type="list", formula1='"' + ",".join(opts) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        L = get_column_letter(cols.index(col) + 1)
        dv.add(f"{L}2:{L}{last}")
    ins = wb.create_sheet("Instructions")
    text = [
        "Review instructions",
        "Each row is ONE event: one pair of animals, one contact type, one time interval. Fill in 'verdict' per row.",
        "Several rows can share a clip (same scene_id): the clip shows every contact of that moment, and",
        "clip_from / clip_to say where in the clip (slow-motion time) this row's event is active.",
        "The two animals of a row are circled in their overlay colours (R1 red, R2 green, R3 blue, R4 cyan,",
        "R5 magenta, R6 yellow) and named in the banner while the event is active (RED border = some label active).",
        "in_sample = FALSE marks events listed only because they overlap a sampled one; partial = TRUE means",
        "the event continues in the previous or next clip (clips never repeat footage).",
        "start_time / end_time are recording times; clip_from / clip_to are positions inside the clip"
        + (f", which plays at {speed_note(speed)}" if speed_note(speed) else "")
        + " — scrub the player to clip_from.",
        "",
        "real: the physical/social contact of the given type is clearly visible for most of the event interval.",
        "flicker: label active but no real contact, a momentary brush (< ~0.3 s) or keypoint jitter.",
        "wrong_type: a real interaction, but of another type; fill in 'correct_type' (N2N, N2AG, N2B, SBS, FOL or NC).",
        "tracking_error: identity swap, merged masks or keypoints on the wrong animal.",
        "unsure: cannot be decided from the clip; excluded from precision.",
        "",
        "NC rows (sample_id N01.1, contact_type NC) are windows in which the pipeline reported NO contact for the",
        "pair. The clip shows only that window and the two animals of the pair are circled. Is there any contact",
        "between them, whatever its type?",
        "nc_correct: no contact between the two animals in the window (the empty label is right).",
        "missed_contact: a real contact is visible although nothing was labelled; put its type in 'correct_type'",
        "(N2N, N2AG, N2B, SBS or FOL).",
        "unsure works for NC rows too and is excluded from the miss rate.",
        "The 'verdict' list is the same for every row: real / flicker / wrong_type / tracking_error belong to",
        "contact rows, nc_correct / missed_contact to NC rows.",
        "",
        "Investigator and boundaries (answer for every row whose verdict is 'real'; both lists are yes / no / na):",
        "investigator_label is the animal the pipeline names as the one doing the sniffing (N2AG, N2B) or the",
        "following (FOL); index.html and the clip banner show it as R2 -> R3 (investigator -> other animal).",
        "investigator_ok: yes = the highlighted animal doing the sniffing / following is the one named in",
        "investigator_label; no = it is the other animal (write that in 'notes'); na = nothing to judge.",
        "Rows without an investigator (N2N, SBS, NC windows) are pre-filled with na.",
        "boundaries_ok: yes = the event starts and ends within about 0.2 s of where you would put them;",
        "no = clearly too early / late / long / short (say which in 'notes'); na = cannot be judged.",
        "NC windows have no boundaries and are pre-filled with na.",
        "",
        "Also fill in 'reviewer' (your initials) and 'notes' (optional). Do not edit the other columns.",
    ]
    for t in text:
        ins.append([t])
    ins["A1"].font = Font(bold=True, size=13)
    ins.column_dimensions["A"].width = 130
    wb.save(path)


def _event_line(r: Dict[str, Any], speed: float) -> str:
    """One <li> of index.html: clip position first (what the player shows), then the recording time."""
    esc = html.escape
    peak = "n/a" if r["peak_score"] is None else f'{r["peak_score"]:.2f}'
    inv = r.get("investigator_label")
    other = [s for s in str(r["pair_label"]).split("-") if s != inv]
    arrow = f' ({esc(str(inv))} &rarr; {esc(other[0])})' if inv and len(other) == 1 else ""
    note = speed_note(speed)
    clip_pos = (f'clip {esc(r["clip_from"])}&ndash;{esc(r["clip_to"])}' + (f" ({note})" if note else "")
                + " &middot; ") if r.get("clip_from") else ""
    tertile = f', {esc(str(r["duration_tertile"]))}' if r.get("duration_tertile") else ""
    flag = "" if r.get("in_sample", True) else " <i>(extra)</i>"
    return (f'<li><b>{esc(r["sample_id"])}</b> <span class="t">{esc(r["contact_type"])}</span> '
            f'{esc(r["pair_label"])}{arrow} &middot; {clip_pos}'
            f'video {esc(r["start_time"])}&ndash;{esc(r["end_time"])} &middot; '
            f'{r["duration_sec"]:.2f} s{tertile} &middot; peak {peak}{flag}</li>')


def write_index_html(rows: List[Dict[str, Any]], path: Path, title: str, speed: float = 1.0) -> None:
    """Self-contained gallery: one card per clip (scene) listing its events, with type filters.

    Each event line gives the clip position first (what the player shows) and then the recording
    ("video") time; `speed` is the clip playback speed.
    """
    esc = html.escape
    types = sorted({r["contact_type"] for r in rows})
    groups: Dict[str, List[Dict[str, Any]]] = {}
    for r in rows:
        groups.setdefault(r["clip"], []).append(r)
    cards = []
    for clip, rs in groups.items():
        head = rs[0].get("scene_id") or rs[0]["sample_id"]
        items = [_event_line(r, speed) for r in rs]
        tset = " ".join(sorted({r["contact_type"] for r in rs}))
        cards.append(f'<div class="card" data-type="{esc(tset)}"><h3>{esc(head)} '
                     f'<span class="n">{len(rs)} event(s)</span></h3>'
                     f'<video controls loop muted preload="metadata" src="{esc(clip)}"></video>'
                     f'<ul>{"".join(items)}</ul></div>')
    time_note = ((f"Clips play at {speed_note(speed)}: " if speed_note(speed) else "")
                 + "'clip' times are what the player shows, 'video' times are the original recording times.")
    btns = '<button data-f="all" class="on">all</button>' + "".join(
        f'<button data-f="{esc(t)}">{esc(t)}</button>' for t in types)
    doc = f"""<!doctype html><html><head><meta charset="utf-8"><title>{esc(title)}</title>
<meta name="viewport" content="width=device-width,initial-scale=1"><style>
:root{{--bg:#fff;--fg:#1a1a1a;--card:#f3f4f6;--acc:#305496}}
@media(prefers-color-scheme:dark){{:root{{--bg:#151515;--fg:#eee;--card:#242424;--acc:#8fb0ee}}}}
body{{background:var(--bg);color:var(--fg);font:15px system-ui,sans-serif;margin:16px}}
#g{{display:grid;grid-template-columns:repeat(auto-fill,minmax(420px,1fr));gap:14px}}
.card{{background:var(--card);border-radius:8px;padding:10px}}.card video{{width:100%}}
.card h3{{margin:0 0 6px}}.t{{color:var(--acc);font-weight:600}}.n{{font-weight:normal;font-size:13px;opacity:.7}}
ul{{margin:6px 0 0;padding-left:18px;font-size:13px}}li{{margin:2px 0}}
button{{margin:0 6px 10px 0;padding:6px 12px;border:1px solid var(--acc);background:none;color:var(--fg);border-radius:6px;cursor:pointer}}
button.on{{background:var(--acc);color:var(--bg)}}</style></head><body>
<h1>{esc(title)}</h1><p>{len(rows)} events in {len(groups)} clips. Red border = a label is active, grey = context.
Each listed event is circled in its animals' colours while active. <b>R2 &rarr; R3</b> = R2 is the investigator
(the animal sniffing or following).</p><p>{html.escape(time_note, quote=False)}</p>
<div>{btns}</div><div id="g">{''.join(cards)}</div><script>
document.querySelectorAll('button').forEach(b=>b.onclick=()=>{{
document.querySelectorAll('button').forEach(x=>x.classList.toggle('on',x===b));
document.querySelectorAll('.card').forEach(c=>c.style.display=(b.dataset.f==='all'||c.dataset.type.split(' ').includes(b.dataset.f))?'':'none');}});
</script></body></html>"""
    path.write_text(doc, encoding="utf-8")


def _event_row(ev: pd.Series, pf: Optional[pd.DataFrame], fps: float) -> Dict[str, Any]:
    """Workbook row shared by both modes (event columns + per-frame aggregates)."""
    st = event_stats(ev, pf)
    dur = ev["duration_sec"] if pd.notna(ev["duration_sec"]) else ev["n_frames"] / fps
    peak = None if pd.isna(ev["peak_score"]) else float(ev["peak_score"])
    return {
        "sample_id": ev["sample_id"], "pair_label": pair_label(ev["pair_key"]),
        "contact_type": ev["contact_type"], "contact_label": TYPE_NAMES.get(ev["contact_type"], ev["contact_type"]),
        "start_time": fmt_time(ev["start_frame"] / fps), "end_time": fmt_time((ev["end_frame"] + 1) / fps),
        "start_frame": int(ev["start_frame"]), "end_frame": int(ev["end_frame"]),
        "duration_sec": round(float(dur), 3), "duration_frames": int(ev["n_frames"]),
        "duration_tertile": ev["duration_tertile"], "peak_score": peak,
        **review_prefill(ev),
        "mean_score": st.get("mean_score"), "secondary_type": st.get("secondary_type"),
        "secondary_score": st.get("secondary_score"),
        **{f"frac_{f}": st.get(f"frac_{f}") for f in QUALITY_FLAGS},
        "mean_nose_nose_dist_bl": st.get("mean_nose_nose_dist_bl"),
        "mean_mask_contact_bl": st.get("mean_mask_contact_bl"),
    }


def cmd_sample(args: argparse.Namespace, overrides: List[str]) -> int:
    """Run the `sample` subcommand."""
    import cv2
    cfg = load_config(overrides)
    for k in ("per_type", "pad_sec", "speed", "seed", "source", "types", "group_by", "nc_windows"):
        v = getattr(args, k, None)
        if v is not None:
            cfg[k] = v
    run = Path(args.run_dir).resolve()
    out = Path(args.out) if args.out else run / "validation"
    (out / "clips").mkdir(parents=True, exist_ok=True)
    (out / "thumbs").mkdir(exist_ok=True)

    events, src = load_events(run, cfg["source"])
    pf = load_per_frame(run)
    if pf is None:
        logger.warning("No per-frame table; clips will lack live values")
    video = find_overlay(run)
    fps = find_fps(run, video)
    logger.info("Source=%s, %d events, video=%s, fps=%.2f", src, len(events), video.name, fps)
    if cfg["types"]:
        keep = {t.strip().upper() for t in str(cfg["types"]).split(",")}
        events = events[events["contact_type"].isin(keep)]
    events_all = events  # every contact event of the run: NC windows must not overlap any of them
    events = events.reset_index(drop=True)
    events["duration_tertile"] = ""
    for t, g in events.groupby("contact_type"):
        events.loc[g.index, "duration_tertile"] = assign_tertiles(g["duration_sec"]).values

    parts, avail, sampled = [], {}, {}
    for t, g in events.groupby("contact_type"):
        s = stratified_sample(g, int(cfg["per_type"]), int(cfg["seed"]))
        avail[t], sampled[t] = len(g), len(s)
        parts.append(s)
    if not parts:
        logger.error("No events to sample")
        return 1
    sample = pd.concat(parts).sort_values(["contact_type", "start_frame"])

    cap = cv2.VideoCapture(str(video))
    nframes = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
    pad = int(round(float(cfg["pad_sec"]) * fps))
    speed, crf = float(cfg["speed"]), int(cfg["crf"])
    colors = overlay_colors(run)
    pf_pairs = index_per_frame(pf)

    rows: List[Dict[str, Any]] = []
    n_clips = 0
    if str(cfg["group_by"]) == "event":
        sample = sample.reset_index(drop=True)
        sample["sample_id"] = [f"V{i + 1:03d}" for i in range(len(sample))]
        for _, ev in sample.iterrows():
            sid = ev["sample_id"]
            clip, sheet = f"clips/{sid}_{ev['contact_type']}_{pair_label(ev['pair_key'])}.mp4", f"thumbs/{sid}.png"
            n = cut_clip(cap, sid, ev, pf_pairs.get(str(ev["pair_key"])), fps, pad, nframes,
                         out / clip, out / sheet, speed, crf, colors)
            a = max(0, int(ev["start_frame"]) - pad)             # first frame of the clip (see cut_clip)
            rows.append({**_event_row(ev, pf, fps), "clip": clip, "thumbnail": sheet,
                         "clip_from": fmt_time((int(ev["start_frame"]) - a) / (fps * speed)),
                         "clip_to": fmt_time((int(ev["end_frame"]) + 1 - a) / (fps * speed)),
                         "frames_written": n})
            n_clips += 1
            logger.info("%s %s %s -> %d frames", sid, ev["contact_type"], pair_label(ev["pair_key"]), n)
    else:
        max_frames = int(round(float(cfg["max_scene_sec"]) * fps)) if cfg.get("max_scene_sec") else None
        scenes = build_scenes(sample, events, pad, nframes, max_frames)
        ids: Dict[Tuple[str, str, int], str] = {}
        for sc in scenes:
            sc["events"]["sample_id"] = [f"{sc['scene_id']}.{i + 1}" for i in range(len(sc["events"]))]
            ids.update(zip(_event_keys(sc["events"]), sc["events"]["sample_id"]))
        for sc in scenes:
            evs, vis = sc["events"], sc["visible"]
            vis["sample_id"] = [ids.get(k, "") for k in _event_keys(vis)]
            clip, sheet = f"clips/{sc['scene_id']}.mp4", f"thumbs/{sc['scene_id']}.png"
            n = cut_scene_clip(cap, sc, vis, pf_pairs, fps, out / clip, out / sheet, speed, crf, colors)
            n_clips += 1
            for _, ev in evs.iterrows():
                lo = max(int(ev["start_frame"]), sc["a"]) - sc["a"]
                hi = min(int(ev["end_frame"]), sc["b"]) + 1 - sc["a"]
                rows.append({**_event_row(ev, pf, fps), "scene_id": sc["scene_id"], "clip": clip,
                             "thumbnail": sheet, "in_sample": bool(ev["in_sample"]),
                             "partial": bool(ev["partial"]),
                             "clip_from": fmt_time(lo / (fps * speed)), "clip_to": fmt_time(hi / (fps * speed)),
                             "frames_written": n})
            logger.info("%s: %d event(s) to review (%d visible), frames %d-%d -> %d frames written",
                        sc["scene_id"], len(evs), len(vis), sc["a"], sc["b"], n)

    n_nc_req, n_nc = int(cfg.get("nc_windows") or 0), 0
    if n_nc_req > 0:
        spans = load_nc_spans(run, cfg["source"])
        if spans.empty:
            logger.warning("nc_windows=%d requested but the events table has no NC rows; skipping", n_nc_req)
        else:
            wins = sample_nc_windows(spans, events_all, pf_pairs, fps, n_nc_req, float(cfg["nc_window_sec"]),
                                     2 * float(cfg["pad_sec"]), int(cfg["seed"]), float(cfg["nc_near_bl"]), nframes)
            if len(wins) < n_nc_req:
                logger.warning("Only %d NC windows available (%d requested)", len(wins), n_nc_req)
            for k in range(len(wins)):
                sid = f"N{k + 1:02d}"
                evdf = wins.iloc[[k]].copy()
                evdf["sample_id"] = f"{sid}.1"
                scene = {"a": int(evdf.iloc[0]["start_frame"]), "b": int(evdf.iloc[0]["end_frame"]),
                         "events": evdf, "visible": evdf, "scene_id": sid, "kind": "nc"}
                clip, sheet = f"clips/{sid}.mp4", f"thumbs/{sid}.png"
                n = cut_scene_clip(cap, scene, evdf, pf_pairs, fps, out / clip, out / sheet, speed, crf, colors)
                n_clips += 1
                ev = evdf.iloc[0]
                near = None if pd.isna(ev["nc_near"]) else bool(ev["nc_near"])
                mind = None if pd.isna(ev["min_centroid_dist_bl"]) else round(float(ev["min_centroid_dist_bl"]), 3)
                rows.append({**_event_row(ev, pf, fps), "scene_id": sid, "clip": clip, "thumbnail": sheet,
                             "in_sample": True, "partial": False, "nc_near": near, "min_centroid_dist_bl": mind,
                             "clip_from": fmt_time(0), "clip_to": fmt_time(n / (fps * speed)),
                             "frames_written": n})
                logger.info("%s NC window %s frames %d-%d (near=%s) -> %d frames", ev["sample_id"],
                            pair_label(ev["pair_key"]), ev["start_frame"], ev["end_frame"], near, n)
                n_nc += 1
    cap.release()

    write_review_xlsx(rows, out / "review.xlsx", speed)
    write_index_html(rows, out / "index.html", f"Contact validation - {run.name}", speed)
    manifest: Dict[str, Any] = {
        "source": src, "run_dir": str(run), "video": str(video), "fps": fps, "params": cfg,
        "seed": cfg["seed"], "available_per_type": avail, "sampled_per_type": sampled,
        "nc_windows_requested": n_nc_req, "nc_windows_sampled": n_nc,
        "n_samples": len(rows), "n_clips": n_clips}
    prov = run / "provenance.json"
    if prov.exists():
        try:
            manifest["provenance"] = json.loads(prov.read_text(encoding="utf-8"))
        except (OSError, ValueError) as exc:
            logger.warning("Could not read %s: %s", prov, exc)
    (out / "manifest.json").write_text(json.dumps(manifest, indent=2, default=str))
    logger.info("Wrote %d clips for %d events to %s", n_clips, len(rows), out)
    return 0


# --------------------------------------------------------------------------- score
def _prec(df: pd.DataFrame) -> Dict[str, Any]:
    n = len(df)
    n_real = int((df["verdict"] == "real").sum())
    denom = int((df["verdict"] != "unsure").sum())
    lo, hi = wilson_ci(n_real, denom)
    return {"n_reviewed": n, "n_real": n_real, "n_unsure": n - denom, "n_decided": denom,
            "precision": n_real / denom if denom else None,
            "ci_low": None if math.isnan(lo) else lo, "ci_high": None if math.isnan(hi) else hi,
            "verdicts": {v: int((df["verdict"] == v).sum()) for v in VERDICTS}}


def _recall_proxy(nc: pd.DataFrame) -> Dict[str, Any]:
    """Miss rate over reviewed no-contact windows: missed_contact / (windows - unsure), Wilson CI.

    It is a rate per sampled window (windows near another animal are drawn first), a proxy
    for recall and not recall itself.
    """
    verdict = nc["verdict"]
    n_unsure = int((verdict == "unsure").sum())
    missed = nc[verdict == "missed_contact"]
    n_decided = len(nc) - n_unsure
    lo, hi = wilson_ci(len(missed), n_decided)
    by_type: Dict[str, int] = {}
    ctypes = missed["correct_type"] if "correct_type" in missed else pd.Series([None] * len(missed))
    for ct in ctypes:
        c = "UNSPECIFIED" if pd.isna(ct) or not str(ct).strip() else str(ct).strip().upper()
        by_type[c] = by_type.get(c, 0) + 1
    n_near = None
    if "nc_near" in nc:
        n_near = int(nc["nc_near"].astype(str).str.strip().str.lower().isin(["true", "1", "1.0"]).sum())
    return {"n_nc_windows": len(nc), "n_nc_correct": int((verdict == "nc_correct").sum()),
            "n_missed": len(missed), "n_unsure": n_unsure, "n_decided": n_decided,
            "miss_rate": len(missed) / n_decided if n_decided else None,
            "ci_low": None if math.isnan(lo) else lo, "ci_high": None if math.isnan(hi) else hi,
            "missed_by_type": by_type, "n_near_windows": n_near}


def _yes_no_stats(sub: pd.DataFrame, col: str) -> Dict[str, Any]:
    """yes / (yes + no) of column `col` with a Wilson CI; 'na' and blank rows are counted apart."""
    v = sub[col].astype(str).str.strip().str.lower()
    n_yes, n_no, n_na = int((v == "yes").sum()), int((v == "no").sum()), int((v == "na").sum())
    n = n_yes + n_no
    lo, hi = wilson_ci(n_yes, n)
    return {"n_scored": n, "n_yes": n_yes, "n_no": n_no, "n_na": n_na,
            "accuracy": n_yes / n if n else None,
            "ci_low": None if math.isnan(lo) else lo, "ci_high": None if math.isnan(hi) else hi}


def review_accuracy(df: pd.DataFrame, col: str, types: Optional[Sequence[str]] = None) -> Optional[Dict[str, Any]]:
    """Accuracy of a yes/no/na review column (investigator_ok, boundaries_ok), overall and per type.

    Only rows with verdict `real` count (a flicker has no investigator or boundaries to judge) and
    only the answers yes and no: accuracy = yes / (yes + no), Wilson 95% CI. `types` restricts the
    rows (investigator: N2AG, N2B, FOL). None when the column is missing (an older review.xlsx) or
    holds no yes / no.
    """
    if col not in df.columns:
        return None
    real = df[df["verdict"] == "real"]
    if types is not None:
        real = real[real["contact_type"].isin(list(types))]
    overall = _yes_no_stats(real, col)
    if overall["n_scored"] == 0:
        return None
    per_type = {t: _yes_no_stats(g, col) for t, g in real.groupby("contact_type")}
    return {"overall": overall, "per_type": {t: st for t, st in per_type.items() if st["n_scored"]}}


def pair_key_of_label(label: Any) -> Optional[str]:
    """'R2-R3' -> '1_2' (the inverse of pair_label); None if it does not parse."""
    try:
        a, b = str(label).split("-")
        return f"{int(a.strip().lstrip('Rr')) - 1}_{int(b.strip().lstrip('Rr')) - 1}"
    except ValueError:
        return None


def _num(x: Any) -> float:
    try:
        return float(x)
    except (TypeError, ValueError):
        return float("nan")


def _row_frames(r: pd.Series, fps: float) -> Optional[Tuple[int, int]]:
    """First and last frame of a review row: start_frame / end_frame when the file has them, else
    derived from start_time (0.1 s resolution) and duration_frames (older review files)."""
    a, b = _num(r.get("start_frame")), _num(r.get("end_frame"))
    if not (math.isnan(a) or math.isnan(b)):
        return int(a), int(b)
    t0, n = parse_time(r.get("start_time")), _num(r.get("duration_frames"))
    if t0 is None or math.isnan(n):
        return None
    a0 = int(round(t0 * fps))
    return a0, a0 + int(n) - 1


def _bout_status(rows: List[pd.Series]) -> str:
    """'real' / 'false' / 'incomplete' from the review rows mapped to one investigation bout.

    false: any event is flicker or tracking_error, or wrong_type with a correct_type outside
    N2N / N2AG / N2B (SBS, FOL, NC). real: every event is real, or wrong_type with a correct_type
    inside N2N / N2AG / N2B (a sub-type confusion does not make the investigation false).
    incomplete: no event fails but one lacks a verdict (blank, unsure, wrong_type without a
    correct_type, or a verdict that does not fit a contact row).
    """
    incomplete = False
    for r in rows:
        v = "" if pd.isna(r.get("verdict")) else str(r["verdict"]).strip().lower()
        ct = "" if pd.isna(r.get("correct_type")) else str(r["correct_type"]).strip().upper()
        if v in ("flicker", "tracking_error"):
            return "false"
        if v == "wrong_type":
            if ct in ("SBS", "FOL", "NC"):
                return "false"
            incomplete = incomplete or ct not in INVESTIGATION_TYPES
        elif v != "real":
            incomplete = True
    return "incomplete" if incomplete else "real"


def bout_precision(review: pd.DataFrame, bouts: pd.DataFrame, fps: float) -> Dict[str, Any]:
    """Precision of the investigation bouts from the event-level verdicts.

    Every N2N / N2AG / N2B row of the review table (all rows, with or without a verdict, sampled or
    extra) is mapped to the bout of the same pair whose frame range overlaps it most; a bout's
    status comes from its rows (see _bout_status). Bouts without any review row were not sampled
    and are not counted. precision = real / (real + false), Wilson 95% CI; incomplete bouts are
    excluded.
    """
    ct = review["contact_type"].astype(str).str.strip().str.upper()
    rows = review[ct.isin(INVESTIGATION_TYPES)]
    by_pair = {str(k): g.reset_index(drop=True) for k, g in bouts.groupby(bouts["pair_key"].astype(str))}
    mapped: Dict[Any, List[pd.Series]] = {}
    meta: Dict[Any, pd.Series] = {}
    n_unmapped = 0
    for _, r in rows.iterrows():
        pk, fr = pair_key_of_label(r.get("pair_label")), _row_frames(r, fps)
        g = by_pair.get(pk) if pk is not None else None
        if g is None or fr is None:
            n_unmapped += 1
            continue
        ov = (np.minimum(g["end_frame"].to_numpy(dtype=float), fr[1])
              - np.maximum(g["start_frame"].to_numpy(dtype=float), fr[0]) + 1)
        k = int(np.argmax(ov))
        if ov[k] <= 0:
            n_unmapped += 1
            continue
        bid = (str(pk), int(g.iloc[k]["start_frame"]))
        mapped.setdefault(bid, []).append(r)
        meta[bid] = g.iloc[k]
    if n_unmapped:
        logger.warning("%d N2N/N2AG/N2B review row(s) overlap no investigation bout of their pair; "
                       "are the review file and investigation_bouts.csv from the same run?", n_unmapped)
    out_bouts = []
    for bid in sorted(mapped, key=lambda b: (b[1], b[0])):
        m = meta[bid]
        out_bouts.append({
            "bout_id": None if "bout_id" not in m or pd.isna(m["bout_id"]) else int(m["bout_id"]),
            "pair_label": pair_label(bid[0]), "start_frame": int(m["start_frame"]),
            "end_frame": int(m["end_frame"]), "status": _bout_status(mapped[bid]),
            "sample_ids": [str(r["sample_id"]) for r in mapped[bid]]})
    n_real = sum(b["status"] == "real" for b in out_bouts)
    n_false = sum(b["status"] == "false" for b in out_bouts)
    lo, hi = wilson_ci(n_real, n_real + n_false)
    return {"n_bouts_total": int(len(bouts)), "n_bouts_with_events": len(out_bouts),
            "n_bouts_scored": n_real + n_false, "n_real": n_real, "n_false": n_false,
            "n_incomplete": len(out_bouts) - n_real - n_false,
            "precision": n_real / (n_real + n_false) if n_real + n_false else None,
            "ci_low": None if math.isnan(lo) else lo, "ci_high": None if math.isnan(hi) else hi,
            "n_rows_mapped": int(sum(len(v) for v in mapped.values())), "n_rows_unmapped": n_unmapped,
            "bouts": out_bouts}


def identity_check(df: pd.DataFrame) -> Optional[Dict[str, Any]]:
    """Summary of a filled identity table (columns time, animal, identity_ok); None without verdicts.

    A checkpoint (one timestamp) is ok when every reviewed animal at it is `yes`. The animal
    figures count (checkpoint, animal) rows. first_failure_time_per_animal is the first `no`.
    """
    if not {"time", "animal", "identity_ok"} <= set(df.columns):
        return None
    d = df.copy()
    d["identity_ok"] = d["identity_ok"].astype(str).str.strip().str.lower()
    d = d[d["identity_ok"].isin(IDENTITY_VERDICTS)]
    if d.empty:
        return None
    d["time"] = d["time"].astype(str).str.strip()
    d["sec"] = d["time"].map(parse_time)
    ok_by_time = d.groupby("time")["identity_ok"].apply(lambda v: bool((v == "yes").all()))
    first_fail: Dict[str, str] = {}
    for animal, g in d[d["identity_ok"] == "no"].groupby("animal"):
        first_fail[str(animal)] = str(g.sort_values("sec").iloc[0]["time"])
    return {"n_checkpoints": int(len(ok_by_time)), "n_checkpoints_ok": int(ok_by_time.sum()),
            "n_animals": int(d["animal"].nunique()), "n_checked": int(len(d)),
            "n_ok": int((d["identity_ok"] == "yes").sum()), "n_no": int((d["identity_ok"] == "no").sum()),
            "n_unsure": int((d["identity_ok"] == "unsure").sum()),
            "first_failure_time_per_animal": dict(sorted(first_fail.items()))}


def compute_summary(df: pd.DataFrame, cfg: dict) -> Dict[str, Any]:
    """Aggregate a filled review table into the summary dict.

    Rows with contact_type NC (no-contact windows) feed only `recall_proxy`; every other
    figure is computed on the contact rows.
    """
    df = df.copy()
    df["verdict"] = df["verdict"].astype(str).str.strip().str.lower()
    is_nc = df["contact_type"].astype(str).str.strip().str.upper() == "NC"
    nc = df[is_nc]
    df = df[~is_nc]
    bad = int((~nc["verdict"].isin(NC_VERDICTS)).sum()) + int((~df["verdict"].isin(VERDICTS)).sum())
    if bad:
        logger.warning("%d review row(s) ignored: verdict not valid for the row type "
                       "(contact rows: %s; NC rows: %s)", bad, ", ".join(VERDICTS), ", ".join(NC_VERDICTS))
    nc = nc[nc["verdict"].isin(NC_VERDICTS)]
    df = df[df["verdict"].isin(VERDICTS)]
    df["duration_sec"] = pd.to_numeric(df["duration_sec"], errors="coerce")
    df["peak_score"] = pd.to_numeric(df.get("peak_score"), errors="coerce")
    # Precision is estimated on the stratified sample only. Events listed just because
    # they overlap a sampled one (in_sample = False, scene mode) over-represent busy
    # moments, so they are summarised apart and never enter the per-type figures.
    extras = None
    if "in_sample" in df.columns:
        core_mask = df["in_sample"].astype(str).str.strip().str.lower().isin(["true", "1", "yes"])
        extras = df[~core_mask]
        df = df[core_mask]
    s: Dict[str, Any] = {"n_reviewed": len(df), "overall": _prec(df) if len(df) else {}}
    s["recall_proxy"] = _recall_proxy(nc) if len(nc) else None
    # Who investigates and where the events start / end, judged on the real events of the sample
    # (None for an older review.xlsx without the columns). bout_precision and identity_check need
    # files next to the review and are added by cmd_score.
    s["investigator_accuracy"] = review_accuracy(df, "investigator_ok", INVESTIGATOR_TYPES)
    s["boundary_accuracy"] = review_accuracy(df, "boundaries_ok")
    s["bout_precision"] = None
    s["identity_check"] = None
    s["per_type"] = {t: _prec(g) for t, g in df.groupby("contact_type")}
    s["extras"] = {"n_reviewed": int(len(extras)), "overall": _prec(extras) if len(extras) else {},
                   "per_type": {t: _prec(g) for t, g in extras.groupby("contact_type")}} if extras is not None else None

    def flick(g: pd.DataFrame) -> Dict[str, Any]:
        return {"n": len(g), "n_flicker": int((g["verdict"] == "flicker").sum()),
                "flicker_rate": float((g["verdict"] == "flicker").mean()) if len(g) else None}
    s["flicker_by_tertile"] = {t: flick(g) for t, g in df.groupby("duration_tertile")} \
        if "duration_tertile" in df else {}
    edges = [0.0] + list(cfg["duration_bins"]) + [float("inf")]
    labels = [f"<{edges[1]:g}s"] + [f"{edges[i]:g}-{edges[i + 1]:g}s" for i in range(1, len(edges) - 2)] \
        + [f">{edges[-2]:g}s"]
    bins = pd.cut(df["duration_sec"], edges, labels=labels, right=False)
    s["flicker_by_duration_bin"] = {str(lab): flick(df[bins == lab]) for lab in labels}
    conf: Dict[str, Dict[str, int]] = {}
    for _, r in df[df["verdict"] == "wrong_type"].iterrows():
        ct = r.get("correct_type")
        c = "UNSPECIFIED" if pd.isna(ct) else str(ct).upper()
        conf.setdefault(r["contact_type"], {}).setdefault(c, 0)
        conf[r["contact_type"]][c] += 1
    s["wrong_type_confusion"] = conf
    s["tracking_error_rate"] = float((df["verdict"] == "tracking_error").mean()) if len(df) else None
    fl = df[df["verdict"] == "flicker"][["sample_id", "contact_type", "duration_sec", "peak_score"]]
    s["flicker_samples"] = [{k: (None if pd.isna(v) else v) for k, v in r.items()}
                            for r in fl.to_dict("records")]
    dec = df[df["verdict"] != "unsure"]
    n_real_all = max(int((dec["verdict"] == "real").sum()), 1)

    def grid(col: str, values: List[float]) -> List[Dict[str, Any]]:
        out = []
        for v in values:
            k = dec[dec[col] >= v]
            n_real = int((k["verdict"] == "real").sum())
            lo, hi = wilson_ci(n_real, len(k))
            out.append({"threshold": v, "n_kept": len(k),
                        "precision": n_real / len(k) if len(k) else None,
                        "ci_low": None if math.isnan(lo) else lo, "ci_high": None if math.isnan(hi) else hi,
                        "recall_of_reviewed_real": n_real / n_real_all})
        return out
    s["min_duration_grid"] = grid("duration_sec", list(cfg["duration_grid"]))
    s["min_peak_score_grid"] = grid("peak_score", list(cfg["score_grid"]))
    return s


def _pct(x: Optional[float]) -> str:
    return "n/a" if x is None else f"{100 * x:.0f}%"


def _render_accuracy_md(s: Dict[str, Any]) -> List[str]:
    """Markdown block for investigator and boundary accuracy ([] when neither was reviewed)."""
    blocks = [(name, s.get(key)) for name, key in (("investigator", "investigator_accuracy"),
                                                    ("boundaries", "boundary_accuracy"))]
    blocks = [(n, b) for n, b in blocks if b]
    if not blocks:
        return []
    L = ["", "## Investigator and event boundaries", "",
         "Judged on events with verdict `real` only, answers yes / no (na and blanks excluded). "
         "Investigator: the animal named in `investigator_label` is the one sniffing (N2AG, N2B) or following (FOL). "
         "Boundaries: start and end within about 0.2 s of where the reviewer would put them. "
         "Accuracy = yes / (yes + no), Wilson 95% CI.", "",
         "| check | type | n scored | yes | no | accuracy | 95% CI |", "|---|---|---|---|---|---|---|"]
    for name, b in blocks:
        for t, st in [("all", b["overall"])] + list(b["per_type"].items()):
            ci = f"{_pct(st['ci_low'])}-{_pct(st['ci_high'])}" if st["ci_low"] is not None else "n/a"
            L.append(f"| {name} | {t} | {st['n_scored']} | {st['n_yes']} | {st['n_no']} | "
                     f"{_pct(st['accuracy'])} | {ci} |")
    return L


def render_md(s: Dict[str, Any]) -> str:
    """Markdown rendering of the summary."""
    L = ["# Validation summary", "",
         f"Reviewed samples: **{s['n_reviewed']}** (stratified sample only"
         + (f"; {s['extras']['n_reviewed']} extra overlapping events reviewed and reported apart" if s.get("extras") else "")
         + "). Precision = real / (reviewed - unsure), "
         "Wilson 95% CI. Small n means wide intervals - always read n.", "",
         "## Precision per type", "",
         "| type | n | real | unsure | precision | 95% CI | verdicts |", "|---|---|---|---|---|---|---|"]
    for t, p in s["per_type"].items():
        ci = f"{_pct(p['ci_low'])}-{_pct(p['ci_high'])}" if p["ci_low"] is not None else "n/a"
        L.append(f"| {t} | {p['n_reviewed']} | {p['n_real']} | {p['n_unsure']} | {_pct(p['precision'])} | {ci} | "
                 + ", ".join(f"{k}={v}" for k, v in p["verdicts"].items() if v) + " |")
    L += ["", "## Flicker rate by duration tertile", "", "| tertile | n | flicker | rate |", "|---|---|---|---|"]
    L += [f"| {k} | {v['n']} | {v['n_flicker']} | {_pct(v['flicker_rate'])} |"
          for k, v in s["flicker_by_tertile"].items()]
    L += ["", "## Flicker rate by duration bin", "", "| bin | n | flicker | rate |", "|---|---|---|---|"]
    L += [f"| {k} | {v['n']} | {v['n_flicker']} | {_pct(v['flicker_rate'])} |"
          for k, v in s["flicker_by_duration_bin"].items()]
    L += ["", "## Wrong-type confusion (label -> correct_type)", ""]
    if s["wrong_type_confusion"]:
        L += [f"- {t}: " + ", ".join(f"{c} x{n}" for c, n in d.items())
              for t, d in s["wrong_type_confusion"].items()]
    else:
        L.append("None.")
    L += ["", f"Tracking-error rate: {_pct(s['tracking_error_rate'])}"]
    rp = s.get("recall_proxy")
    if rp:
        ci = f"{_pct(rp['ci_low'])}-{_pct(rp['ci_high'])}" if rp["ci_low"] is not None else "n/a"
        L += ["", "## Recall proxy: no-contact windows", "",
              "Windows where the pipeline reported no contact for a pair, drawn at random from long NC spans "
              "(pairs that come within the near-distance threshold are drawn first), reviewed for any contact. "
              "Miss rate = missed_contact / (windows - unsure), Wilson 95% CI. It is a rate per sampled window "
              "and over-represents close pairs, so it is a proxy for recall, not recall itself.", "",
              "| windows | nc_correct | missed_contact | unsure | miss rate | 95% CI | near windows |",
              "|---|---|---|---|---|---|---|",
              f"| {rp['n_nc_windows']} | {rp['n_nc_correct']} | {rp['n_missed']} | {rp['n_unsure']} | "
              f"{_pct(rp['miss_rate'])} | {ci} | {'n/a' if rp['n_near_windows'] is None else rp['n_near_windows']} |", ""]
        L.append("Missed contacts by type: " + (", ".join(f"{c} x{n}" for c, n in rp["missed_by_type"].items())
                                                 if rp["missed_by_type"] else "none") + ".")
    L += _render_accuracy_md(s)
    bp = s.get("bout_precision")
    if bp and bp["n_bouts_with_events"]:
        ci = f"{_pct(bp['ci_low'])}-{_pct(bp['ci_high'])}" if bp["ci_low"] is not None else "n/a"
        L += ["", "## Investigation-bout precision", "",
              "N2N, N2AG and N2B events merged per pair (`investigation_bouts.csv`). Every reviewed event is "
              "mapped to the bout of its pair that overlaps it. A bout is *real* when each of its reviewed events "
              "is real, or a type confusion inside N2N / N2AG / N2B; *false* when one is a flicker, a tracking "
              "error or really SBS / FOL / NC; *incomplete* (an event without a usable verdict) is excluded. "
              "Precision = real / (real + false), Wilson 95% CI. Bouts without a reviewed event were not sampled.", "",
              "| bouts with reviewed events | scored | real | false | incomplete | precision | 95% CI |",
              "|---|---|---|---|---|---|---|",
              f"| {bp['n_bouts_with_events']} of {bp['n_bouts_total']} | {bp['n_bouts_scored']} | {bp['n_real']} | "
              f"{bp['n_false']} | {bp['n_incomplete']} | {_pct(bp['precision'])} | {ci} |", ""]
        false = [b for b in bp["bouts"] if b["status"] == "false"]
        L.append("False bouts: " + (", ".join(f"{b['pair_label']} frames {b['start_frame']}-{b['end_frame']} "
                                              f"({'/'.join(b['sample_ids'])})" for b in false) if false else "none") + ".")
    ic = s.get("identity_check")
    if ic:
        L += ["", "## Identity persistence", "",
              f"Identity persistence: {ic['n_checkpoints_ok']}/{ic['n_checkpoints']} checkpoints ok "
              f"(every animal `yes`); {ic['n_ok']}/{ic['n_checked']} animal checks ok, {ic['n_no']} no, "
              f"{ic['n_unsure']} unsure; {ic['n_animals']} animals. "
              + ("; ".join(f"{a} swapped or lost at {t}" for a, t in ic["first_failure_time_per_animal"].items())
                 if ic["first_failure_time_per_animal"] else "No swap or loss reported") + "."]
    L += ["", "## Precision if only events with duration >= d were kept", "",
          "| d (s) | n kept | precision | 95% CI |", "|---|---|---|---|"]
    for g in s["min_duration_grid"]:
        L.append(f"| {g['threshold']:g} | {g['n_kept']} | {_pct(g['precision'])} | "
                 f"{_pct(g['ci_low'])}-{_pct(g['ci_high'])} |")
    L += ["", "## Precision if only events with peak_score >= s were kept", "",
          "| s | n kept | precision | 95% CI |", "|---|---|---|---|"]
    for g in s["min_peak_score_grid"]:
        L.append(f"| {g['threshold']:g} | {g['n_kept']} | {_pct(g['precision'])} | "
                 f"{_pct(g['ci_low'])}-{_pct(g['ci_high'])} |")
    L += ["", "## Samples marked flicker", "", "| sample | type | duration (s) | peak_score |", "|---|---|---|---|"]
    L += [f"| {r['sample_id']} | {r['contact_type']} | {r['duration_sec']} | {r['peak_score']} |"
          for r in s["flicker_samples"]]
    return "\n".join(L) + "\n"


def plot_precision(s: Dict[str, Any], path: Path) -> None:
    """Precision per type with Wilson CI error bars."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    items = [(t, p) for t, p in s["per_type"].items() if p["precision"] is not None]
    fig, ax = plt.subplots(figsize=(6, 4))
    if items:
        x = np.arange(len(items))
        pr = np.array([p["precision"] for _, p in items])
        lo = pr - np.array([p["ci_low"] for _, p in items])
        hi = np.array([p["ci_high"] for _, p in items]) - pr
        ax.bar(x, pr, yerr=[lo, hi], capsize=4, color="#305496")
        ax.set_xticks(x, [f"{t}\n(n={p['n_decided']})" for t, p in items])
    ax.set_ylim(0, 1.05)
    ax.set_ylabel("precision (real / decided)")
    ax.set_title("Contact precision by type (Wilson 95% CI)")
    fig.tight_layout()
    fig.savefig(path, dpi=120)
    plt.close(fig)


# --------------------------------------------------------------------------- inter-rater
POSITIVE_VERDICTS = {"real", "nc_correct"}                                    # the label is right
NEGATIVE_VERDICTS = {"flicker", "wrong_type", "tracking_error", "missed_contact"}  # the label is wrong


def binary_verdict(v: Any) -> Optional[bool]:
    """True = the label is right (real / nc_correct), False = wrong, None = unsure or blank."""
    v = str(v).strip().lower()
    if v in POSITIVE_VERDICTS:
        return True
    if v in NEGATIVE_VERDICTS:
        return False
    return None


def cohen_kappa(a: Sequence[bool], b: Sequence[bool]) -> Dict[str, Any]:
    """Cohen's kappa for two raters on a binary label, computed by hand.

    kappa = (p_o - p_e) / (1 - p_e) with p_o the observed agreement and p_e the agreement
    expected from each rater's marginal rate of True. kappa is None when it is undefined
    (no rows, or p_e = 1 because both raters used one single class).
    The 2x2 table has rater A in the rows and rater B in the columns, in the order
    [True, False].
    """
    n = len(a)
    tt = sum(1 for x, y in zip(a, b) if x and y)
    tf = sum(1 for x, y in zip(a, b) if x and not y)
    ft = sum(1 for x, y in zip(a, b) if not x and y)
    ff = n - tt - tf - ft
    out: Dict[str, Any] = {"n": n, "n_agree": tt + ff, "raw_agreement": None, "p_expected": None, "kappa": None,
                           "table": [[tt, tf], [ft, ff]]}
    if n == 0:
        return out
    p_o = (tt + ff) / n
    p_a, p_b = (tt + tf) / n, (tt + ft) / n
    p_e = p_a * p_b + (1 - p_a) * (1 - p_b)
    out.update(raw_agreement=p_o, p_expected=p_e)
    if p_e < 1:
        out["kappa"] = (p_o - p_e) / (1 - p_e)
    return out


def _agreement(m: pd.DataFrame) -> Dict[str, Any]:
    """Kappa statistics of the matched rows `m` (columns bin_a, bin_b, verdict_a, verdict_b)."""
    st = cohen_kappa([bool(x) for x in m["bin_a"]], [bool(x) for x in m["bin_b"]])
    st["raw_agreement_full_verdict"] = float((m["verdict_a"] == m["verdict_b"]).mean()) if len(m) else None
    return st


def compare_reviews(a: pd.DataFrame, b: pd.DataFrame) -> Dict[str, Any]:
    """Inter-rater agreement of two filled review tables, rows matched by sample_id.

    The verdict is reduced to "the label is right": real -> yes, flicker / wrong_type /
    tracking_error -> no, unsure (or blank) -> row excluded. `overall` covers the contact
    rows; NC-window rows (nc_correct yes, missed_contact no) are a group of their own in
    `per_type` ("NC") and never enter `overall`.
    """
    def prep(d: pd.DataFrame) -> pd.DataFrame:
        d = d[d["verdict"].notna() & d["sample_id"].notna()].copy()
        d["sample_id"] = d["sample_id"].astype(str).str.strip()
        d["verdict"] = d["verdict"].astype(str).str.strip().str.lower()
        return d.drop_duplicates("sample_id")

    a, b = prep(a), prep(b)

    def side(d: pd.DataFrame, suf: str) -> pd.DataFrame:
        keep = [c for c in ("sample_id", "verdict", "contact_type", "pair_label") if c in d]
        return d[keep].rename(columns={c: f"{c}_{suf}" for c in keep if c != "sample_id"})

    m = side(a, "a").merge(side(b, "b"), on="sample_id", how="inner")
    m["contact_type"] = m["contact_type_a"].astype(str).str.strip().str.upper()
    n_mismatch = 0
    for c in ("contact_type", "pair_label"):
        if f"{c}_a" in m and f"{c}_b" in m:
            n_mismatch += int((m[f"{c}_a"].astype(str) != m[f"{c}_b"].astype(str)).sum())
    if n_mismatch:
        logger.warning("%d matched row(s) differ in contact_type / pair_label: are both files from the same "
                       "`sample` run?", n_mismatch)
    m["bin_a"], m["bin_b"] = m["verdict_a"].map(binary_verdict), m["verdict_b"].map(binary_verdict)
    usable = m[m["bin_a"].notna() & m["bin_b"].notna()]
    is_nc = usable["contact_type"] == "NC"
    res: Dict[str, Any] = {
        "n_rows_a": len(a), "n_rows_b": len(b), "n_matched": len(m),
        "n_only_a": len(set(a["sample_id"]) - set(b["sample_id"])),
        "n_only_b": len(set(b["sample_id"]) - set(a["sample_id"])),
        "n_unsure_excluded": len(m) - len(usable), "n_id_mismatch": n_mismatch,
        "overall": _agreement(usable[~is_nc]),
        "per_type": {t: _agreement(g) for t, g in usable.groupby("contact_type")},
        "disagreements": [
            {"sample_id": r.sample_id, "contact_type": r.contact_type, "verdict_a": r.verdict_a,
             "verdict_b": r.verdict_b} for r in usable[usable["bin_a"] != usable["bin_b"]].itertuples()],
    }
    for side, d in (("a", a), ("b", b)):
        names = d["reviewer"].dropna().astype(str).str.strip() if "reviewer" in d else pd.Series(dtype=str)
        res[f"reviewer_{side}"] = names.mode().iloc[0] if len(names) else None
    return res


def render_compare_md(r: Dict[str, Any], name_a: str, name_b: str) -> str:
    """Markdown rendering of compare_reviews()."""
    def k(x: Optional[float]) -> str:
        return "undefined" if x is None else f"{x:.3f}"

    ra, rb = r.get("reviewer_a") or "?", r.get("reviewer_b") or "?"
    L = ["# Inter-rater agreement", "",
         f"Rater A: {ra} ({name_a}); rater B: {rb} ({name_b}). Rows matched by sample_id: **{r['n_matched']}** "
         f"({r['n_unsure_excluded']} excluded because either rater said unsure; {r['n_only_a']} only in A, "
         f"{r['n_only_b']} only in B).", "",
         "The verdict is reduced to *the label is right*: real = yes; flicker, wrong_type and tracking_error = no. "
         "NC-window rows count nc_correct as yes and missed_contact as no and are reported on their own line "
         "(`NC`), not in `overall`. Cohen's kappa = (p_o - p_e) / (1 - p_e), p_e from each rater's own rate of "
         "*yes*; it is undefined when both raters used one class only.", "",
         "| group | n | agree | raw agreement | kappa | same full verdict |", "|---|---|---|---|---|---|"]
    for name, st in [("overall (contact rows)", r["overall"])] + list(r["per_type"].items()):
        L.append(f"| {name} | {st['n']} | {st['n_agree']} | {_pct(st['raw_agreement'])} | {k(st['kappa'])} | "
                 f"{_pct(st['raw_agreement_full_verdict'])} |")
    t = r["overall"]["table"]
    L += ["", "## 2x2 table, overall (contact rows)", "",
          "| | B: real | B: not real |", "|---|---|---|",
          f"| **A: real** | {t[0][0]} | {t[0][1]} |", f"| **A: not real** | {t[1][0]} | {t[1][1]} |", ""]
    L += ["## Disagreements", ""]
    if r["disagreements"]:
        L += ["| sample | type | A | B |", "|---|---|---|---|"]
        L += [f"| {d['sample_id']} | {d['contact_type']} | {d['verdict_a']} | {d['verdict_b']} |"
              for d in r["disagreements"]]
    else:
        L.append("None.")
    return "\n".join(L) + "\n"


def cmd_compare(args: argparse.Namespace) -> int:
    """Run `score --compare A B`: write interrater_summary.{json,md} next to review A."""
    pa, pb = Path(args.compare[0]), Path(args.compare[1])
    a, b = pd.read_excel(pa, sheet_name="review"), pd.read_excel(pb, sheet_name="review")
    r = compare_reviews(a, b)
    if not r["n_matched"]:
        logger.error("No sample_id with a verdict in both %s and %s", pa, pb)
        return 1
    (pa.parent / "interrater_summary.json").write_text(json.dumps(r, indent=2, default=str))
    (pa.parent / "interrater_summary.md").write_text(render_compare_md(r, str(pa), str(pb)), encoding="utf-8")
    ov = r["overall"]
    logger.info("Compared %d matched rows: agreement %s, kappa %s -> %s", r["n_matched"],
                _pct(ov["raw_agreement"]), "undefined" if ov["kappa"] is None else f"{ov['kappa']:.3f}", pa.parent)
    return 0


def resolve_run(review: Path, run_arg: Optional[str] = None) -> Tuple[Path, Dict[str, Any]]:
    """Run directory a review.xlsx belongs to, plus the manifest.json next to it ({} if none).

    Order: --run, the manifest's run_dir when it exists on this machine (a manifest written on
    Bunya names a path that is not there after the download), else the folder above the review's.
    """
    manifest: Dict[str, Any] = {}
    mp = review.parent / "manifest.json"
    if mp.exists():
        try:
            manifest = json.loads(mp.read_text(encoding="utf-8"))
        except (OSError, ValueError) as exc:
            logger.warning("Could not read %s: %s", mp, exc)
    if run_arg:
        return Path(run_arg).resolve(), manifest
    rd = manifest.get("run_dir")
    if rd and Path(str(rd)).is_dir():
        return Path(str(rd)), manifest
    return review.parent.parent, manifest


def load_bout_precision(run: Path, manifest: Dict[str, Any], review: pd.DataFrame) -> Optional[Dict[str, Any]]:
    """bout_precision of a review table against <run>/contacts/investigation_bouts.csv (None, with a
    warning, when the file is missing or unusable)."""
    p = run / "contacts" / "investigation_bouts.csv"
    if not p.exists():
        logger.warning("%s not found: investigation-bout precision not computed "
                       "(re-run postprocess_contacts_simple.py, or pass --run <run_dir>)", p)
        return None
    bouts = pd.read_csv(p)
    if not {"pair_key", "start_frame", "end_frame"} <= set(bouts.columns):
        logger.warning("%s lacks pair_key / start_frame / end_frame: investigation-bout precision not computed", p)
        return None
    fps = float(manifest["fps"]) if isinstance(manifest.get("fps"), (int, float)) else find_fps(run, None)
    return bout_precision(review, bouts, fps)


def find_identity_xlsx(review: Path, run: Path) -> Optional[Path]:
    """identity.xlsx written by `identity`: next to the review (validation/identity/) or under the run."""
    for p in (review.parent / "identity" / "identity.xlsx", run / "validation" / "identity" / "identity.xlsx"):
        if p.exists():
            return p
    return None


def cmd_score(args: argparse.Namespace, overrides: List[str]) -> int:
    """Run the `score` subcommand."""
    cfg = load_config(overrides)
    path = Path(args.review).resolve()
    full = pd.read_excel(path, sheet_name="review")
    df = full[full["verdict"].notna()]
    run, manifest = resolve_run(path, getattr(args, "run", None))
    ident = None
    ip = find_identity_xlsx(path, run)
    if ip is not None:
        ident = identity_check(pd.read_excel(ip, sheet_name="identity"))
    if df.empty and ident is None:
        logger.error("No verdicts filled in %s", path)
        return 1
    s = compute_summary(df, cfg)
    s["bout_precision"] = load_bout_precision(run, manifest, full) if len(df) else None
    s["identity_check"] = ident
    (path.parent / "validation_summary.json").write_text(json.dumps(s, indent=2, default=str))
    (path.parent / "validation_summary.md").write_text(render_md(s), encoding="utf-8")
    plot_precision(s, path.parent / "validation_precision.png")
    logger.info("Scored %d reviewed samples -> %s", s["n_reviewed"], path.parent)
    return 0


# --------------------------------------------------------------------------- identity persistence
def identity_checkpoints(nframes: int, fps: float, every_sec: float, max_n: int) -> List[int]:
    """Checkpoint frames at 0, every_sec, 2 * every_sec, ... (< nframes), at most max_n of them.

    A session with more checkpoints than max_n keeps max_n of them spread evenly over the
    session (always the first and the last one); max_n <= 0 means no cap.
    """
    step = max(float(every_sec) * fps, 1.0)
    frames: List[int] = []
    k = 0
    while int(round(k * step)) < nframes:
        frames.append(int(round(k * step)))
        k += 1
    if max_n and max_n > 0 and len(frames) > max_n:
        pick = np.unique(np.round(np.linspace(0, len(frames) - 1, int(max_n))).astype(int))
        frames = [frames[i] for i in pick]
    return frames


def find_raw_video(run: Path) -> Optional[Path]:
    """Raw video named by video_path in config_used.yaml (as is, then relative to the repository, then
    under data/raw/); None if it is not on this machine."""
    try:
        import yaml
        vp = (yaml.safe_load((run / "config_used.yaml").read_text(encoding="utf-8")) or {}).get("video_path")
    except Exception:  # noqa: BLE001  missing file or bad yaml
        return None
    if not vp:
        return None
    for p in (Path(str(vp)), _PROJECT_ROOT / str(vp), _PROJECT_ROOT / "data" / "raw" / Path(str(vp)).name):
        if p.is_file():
            return p
    return None


def animal_centroids(rows: pd.DataFrame) -> Dict[int, Tuple[int, int]]:
    """{animal slot: (x, y) centroid pixel} of one frame from its per-pair rows.

    Every pair row carries the centroids of both animals, so any pair that includes an animal
    gives its centroid; animals without a valid centroid in the frame are left out.
    """
    out: Dict[int, Tuple[int, int]] = {}
    for _, row in rows.iterrows():
        si, sj = _pair_slots(row["pair_key"])
        for slot, kx, ky in ((si, "centroid_i_x", "centroid_i_y"), (sj, "centroid_j_x", "centroid_j_y")):
            if slot is not None and slot not in out:
                p = _px(row, kx, ky)
                if p is not None:
                    out[slot] = p
    return out


def draw_identity(frame: np.ndarray, centroids: Dict[int, Tuple[int, int]],
                  colors: List[Tuple[int, int, int]], t: float, fidx: int) -> np.ndarray:
    """Dot and large R<n> label at each animal's centroid, plus a time stamp (in place).

    Independent of the labels the overlay video burns in: they come from the per-frame table.
    """
    import cv2
    h = frame.shape[0]
    scale = max(1.0, h / 640.0)
    r = int(round(9 * scale))
    for slot, (x, y) in sorted(centroids.items()):
        col = tuple(int(c) for c in colors[slot % len(colors)])
        cv2.circle(frame, (x, y), r + 2, (0, 0, 0), -1, cv2.LINE_AA)
        cv2.circle(frame, (x, y), r, col, -1, cv2.LINE_AA)
        _put(frame, slot_label(slot), (x + r + 4, y - r - 4), 0.9 * scale, col, max(2, int(round(2 * scale))))
    _put(frame, f"identity check  t={fmt_time(t)}  frame {fidx}", (8, h - 10), 0.6 * scale, (255, 255, 255),
         max(1, int(round(scale))))
    return frame


def write_identity_xlsx(records: List[Dict[str, Any]], slots: List[int], path: Path) -> None:
    """identity.xlsx: one row per (checkpoint, animal) with a yes / no / unsure dropdown, plus Instructions."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    cols = ["time", "frame", "image", "animal", "identity_ok", "notes"]
    wb = Workbook()
    ws = wb.active
    ws.title = "identity"
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
    for rec in records:
        for slot in slots:
            ws.append([rec["time"], rec["frame"], rec["image"], slot_label(slot), None,
                       None if slot in rec["drawn"] else "not drawn: no centroid for this animal in this frame"])
            cell = ws.cell(row=ws.max_row, column=cols.index("image") + 1)
            cell.hyperlink = rec["image"]
            cell.font = Font(color="0563C1", underline="single")
    ws.freeze_panes = "A2"
    for letter, width in (("A", 10), ("B", 9), ("C", 28), ("D", 8), ("E", 13), ("F", 60)):
        ws.column_dimensions[letter].width = width
    dv = DataValidation(type="list", formula1='"' + ",".join(IDENTITY_VERDICTS) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"E2:E{max(ws.max_row, 2)}")
    ins = wb.create_sheet("Instructions")
    text = [
        "Identity check instructions",
        "Each row is ONE animal at ONE checkpoint. Open the image of the row (or identity.html): every animal",
        "carries a large label R1..R6 drawn from the tracker's own centroid, in its overlay colour.",
        "Fill in identity_ok for every row, including the first checkpoint (time 00:00.0), which is the reference.",
        "",
        "yes: the label on this animal is the same individual as at time 0. Compare with the previous checkpoint:",
        "     same coat marks, size, scars or tail, and the animal has not changed places with another one.",
        "no: swapped or lost - the label sits on another individual, or on nothing. Write in 'notes' which",
        "     animal it is now (for example 'R3 is now on the animal that was R5').",
        "unsure: cannot be decided from the image; excluded from the counts of yes / no.",
        "A row 'not drawn' in notes means the tracker had no centroid for the animal in that frame: answer no",
        "if the animal is visible in the image, unsure otherwise.",
        "After the first no for an animal, judge the later checkpoints against time 0 as well, not against the",
        "swapped label.",
        "",
        "Do not edit time, frame, image or animal. 'notes' is optional except for a no.",
    ]
    for t in text:
        ins.append([t])
    ins["A1"].font = Font(bold=True, size=13)
    ins.column_dimensions["A"].width = 120
    wb.save(path)


def write_identity_html(records: List[Dict[str, Any]], slots: List[int],
                        colors: List[Tuple[int, int, int]], path: Path, title: str) -> None:
    """identity.html: a contact sheet, one card per checkpoint (image, time, row of animal labels)."""
    esc = html.escape
    cards = []
    for k, rec in enumerate(records, 1):
        chips = []
        for slot in slots:
            b, g, r = (int(c) for c in colors[slot % len(colors)])
            miss = slot not in rec["drawn"]
            chips.append(f'<span class="chip{" miss" if miss else ""}" style="--c:rgb({r},{g},{b})">'
                         f'{esc(slot_label(slot))}{" (not drawn)" if miss else ""}</span>')
        cards.append(f'<div class="card"><h3>{esc(rec["time"])} <span class="n">frame {rec["frame"]} &middot; '
                     f'checkpoint {k} of {len(records)}</span></h3>'
                     f'<a href="{esc(rec["image"])}"><img loading="lazy" src="{esc(rec["image"])}" '
                     f'alt="identity check at {esc(rec["time"])}"></a><div>{"".join(chips)}</div></div>')
    doc = f"""<!doctype html><html><head><meta charset="utf-8"><title>{esc(title)}</title>
<meta name="viewport" content="width=device-width,initial-scale=1"><style>
:root{{--bg:#fff;--fg:#1a1a1a;--card:#f3f4f6;--acc:#305496}}
@media(prefers-color-scheme:dark){{:root{{--bg:#151515;--fg:#eee;--card:#242424;--acc:#8fb0ee}}}}
body{{background:var(--bg);color:var(--fg);font:15px system-ui,sans-serif;margin:16px}}
#g{{display:grid;grid-template-columns:repeat(auto-fill,minmax(440px,1fr));gap:14px}}
.card{{background:var(--card);border-radius:8px;padding:10px}}.card img{{width:100%;display:block}}
.card h3{{margin:0 0 6px}}.n{{font-weight:normal;font-size:13px;opacity:.7}}
.chip{{display:inline-block;margin:6px 6px 0 0;padding:2px 10px;border-radius:12px;font-weight:600;
border:3px solid var(--c);color:var(--fg)}}.chip.miss{{border-style:dashed;opacity:.6}}</style></head><body>
<h1>{esc(title)}</h1><p>{len(records)} checkpoints. Each animal carries its tracker label (R1..R6) at its centroid.
For every animal ask: is this the same individual as at time 0 (compare with the previous checkpoint)?
Record the answer per animal in <code>identity.xlsx</code>: yes, no (swapped or lost; note who it is now) or unsure.</p>
<div id="g">{''.join(cards)}</div></body></html>"""
    path.write_text(doc, encoding="utf-8")


def cmd_identity(args: argparse.Namespace, overrides: List[str]) -> int:
    """Run the `identity` subcommand: checkpoint images with per-animal labels + identity.xlsx / .html."""
    import cv2
    cfg = load_config(overrides)
    ic = cfg.get("identity") or {}
    every = float(args.every if args.every is not None else ic.get("every_sec", 60))
    max_n = int(args.max if args.max is not None else ic.get("max_frames", 40))
    if every <= 0:
        logger.error("--every must be > 0 (got %s)", every)
        return 1
    run = Path(args.run_dir).resolve()
    out = Path(args.out) if args.out else run / "validation" / "identity"
    pf = load_per_frame(run)
    if pf is None or "pair_key" not in pf.columns:
        logger.error("No per-frame table (contacts_per_frame.csv or the PerFrame sheet of results.xlsx) under %s", run)
        return 1
    try:
        video = find_overlay(run)
    except FileNotFoundError:
        video = find_raw_video(run)
        if video is None:
            logger.error("No overlay video in %s/overlays and no raw video from config_used.yaml", run)
            return 1
        logger.warning("No overlay video; using the raw video %s", video)
    fps = find_fps(run, video)
    cap = cv2.VideoCapture(str(video))
    nframes = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
    frames = identity_checkpoints(nframes, fps, every, max_n)
    if not frames:
        logger.error("Video %s has no frames", video)
        return 1
    out.mkdir(parents=True, exist_ok=True)
    colors = overlay_colors(run)
    pf = pf.assign(frame_idx=pd.to_numeric(pf["frame_idx"], errors="coerce"))
    slots = sorted({s for k in pf["pair_key"].astype(str).unique() for s in _pair_slots(k) if s is not None})
    sel = pf[pf["frame_idx"].isin(frames)]
    by_frame = {int(f): g for f, g in sel.groupby("frame_idx")}
    records: List[Dict[str, Any]] = []
    for f in frames:
        cap.set(cv2.CAP_PROP_POS_FRAMES, f)
        ok, fr = cap.read()
        if not ok:
            logger.warning("Could not read frame %d of %s; checkpoint skipped", f, video.name)
            continue
        t = f / fps
        cen = animal_centroids(by_frame[f]) if f in by_frame else {}
        name = f"identity_{t:07.2f}s.png"
        cv2.imwrite(str(out / name), draw_identity(fr, cen, colors, t, f))
        records.append({"time": fmt_time(t), "frame": f, "image": name, "drawn": sorted(cen)})
        logger.info("t=%s frame %d: %d of %d animals drawn", fmt_time(t), f, len(cen), len(slots))
    cap.release()
    if not records:
        logger.error("No checkpoint frame could be read")
        return 1
    write_identity_xlsx(records, slots, out / "identity.xlsx")
    write_identity_html(records, slots, colors, out / "identity.html", f"Identity check - {run.name}")
    logger.info("Wrote %d checkpoints (every %g s, max %d) for %d animals to %s",
                len(records), every, max_n, len(slots), out)
    return 0


def main(argv: Optional[List[str]] = None) -> int:
    """CLI entry point."""
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    sp = sub.add_parser("sample", help="sample events and build evidence clips")
    sp.add_argument("run_dir")
    sp.add_argument("--per-type", dest="per_type", type=int)
    sp.add_argument("--pad-sec", dest="pad_sec", type=float)
    sp.add_argument("--speed", type=float)
    sp.add_argument("--seed", type=int)
    sp.add_argument("--types")
    sp.add_argument("--source", choices=["events", "bouts"])
    sp.add_argument("--group-by", dest="group_by", choices=["scene", "event"],
                    help="scene (default): one clip per group of overlapping events; event: one clip each")
    sp.add_argument("--nc-per-run", dest="nc_windows", type=int,
                    help="also sample N windows where no contact was reported (recall proxy); default 0 = off")
    sp.add_argument("--out")
    sc = sub.add_parser("score", help="score a filled review.xlsx, or compare two reviewers (--compare)")
    sc.add_argument("review", nargs="?")
    sc.add_argument("--compare", nargs=2, metavar=("REVIEW_A", "REVIEW_B"),
                    help="inter-rater agreement (Cohen's kappa) of two filled review.xlsx files")
    sc.add_argument("--run", help="run directory holding contacts/investigation_bouts.csv (default: from "
                                  "manifest.json next to the review, else the folder above it)")
    idp = sub.add_parser("identity", help="checkpoint images with per-animal labels, to check that identities persist")
    idp.add_argument("run_dir")
    idp.add_argument("--every", type=float, help="seconds between checkpoints (default identity.every_sec = 60)")
    idp.add_argument("--max", type=int, help="most checkpoints, spread evenly (default identity.max_frames = 40)")
    idp.add_argument("--out", help="output folder (default <run>/validation/identity)")
    args, extra = ap.parse_known_args(argv)
    overrides = [e for e in extra if "=" in e]
    if args.cmd == "sample":
        return cmd_sample(args, overrides)
    if args.cmd == "identity":
        return cmd_identity(args, overrides)
    if args.compare and args.review and "=" in args.review and not Path(args.review).exists():
        overrides.append(args.review)  # a key=value override argparse took for the positional
        args.review = None
    if bool(args.review) == bool(args.compare):
        ap.error("score needs either a review.xlsx or --compare REVIEW_A REVIEW_B")
    return cmd_compare(args) if args.compare else cmd_score(args, overrides)


if __name__ == "__main__":
    sys.exit(main())
